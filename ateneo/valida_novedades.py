#!/usr/bin/env python3
"""Revisa un archivo A PROCESAR del Servicio de Novedades antes de cargarlo.

No corrige nada: lista lo que esta mal para que Franco lo arregle a mano.

El archivo trae la hoja que procesa la macro (A1 = "IdCliente", IDs desde la
columna I) y la planilla del vendedor (B5 = "RAZON SOCIAL", IDs desde la J,
clientes desde la fila 6). Puede traer otras hojas auxiliares.

Controles:
  1. IDs: iguales en las dos hojas, sin repetidos, y los esperados del mes
     (novedades + reimpresiones) en el mismo orden; ISBN de la fila 3 contra
     el ID de la fila 5
  2. hoja de proceso == planilla sin la columna RAZON SOCIAL, celda por celda
  3. cantidades: todas numericas (una celda vacia hace que la macro tire
     "No coinciden los tipos")
  4. columnas fijas: tipocomer Q, DirTrans 2, Trans Retira 2, una sola lista
  5. contra el maestro de clientes: el cliente existe, no esta dado de baja,
     el DirEnt es una direccion ENTR, el vendedor coincide y la razon social
     se parece al nombre del maestro
  6. clientes repetidos y fila de cuotas (ultimo cliente + 2)

Uso:
    python valida_novedades.py ARCHIVO.xlsx --maestro Maestro_clientes.xls \
        --ids historial/2026-10/novedades_octubre.csv historial/2026-09/reimpresiones_octubre.csv
"""
import argparse
import difflib
import re
import unicodedata

import pandas as pd
from openpyxl import load_workbook


def norm(txt):
    txt = unicodedata.normalize('NFKD', str(txt)).encode('ascii', 'ignore').decode()
    txt = re.sub(r'\b(S\.?R\.?L|S\.?A\.?S?|S\.?H|SRL|SAS|SA)\b\.?', ' ', txt.upper())
    return re.sub(r'[^A-Z0-9 ]', ' ', re.sub(r'\s+', ' ', txt)).strip()


def parecido(a, b):
    # "BARRIOS (EX 20442)", "AFONSO (EX VERCELLONE)": lo de entre parentesis es historia
    a, b = norm(re.sub(r'\(.*?\)', ' ', str(a))), norm(re.sub(r'\(.*?\)', ' ', str(b)))
    if not a or not b:
        return 0
    if a in b or b in a:
        return 1
    # la planilla suele poner solo el apellido: alcanza con que coincida una palabra larga
    if {w for w in a.split() if len(w) > 3} & {w for w in b.split() if len(w) > 3}:
        return 1
    # los nombres de persona vienen en distinto orden (APELLIDO NOMBRE)
    return max(difflib.SequenceMatcher(None, a, b).ratio(),
               difflib.SequenceMatcher(None, ' '.join(sorted(a.split())),
                                       ' '.join(sorted(b.split()))).ratio())


def buscar_hojas(wb):
    proceso = next((ws for ws in wb.worksheets
                    if str(ws['A1'].value).strip() == 'IdCliente'), None)
    vendedor = next((ws for ws in wb.worksheets
                     if str(ws['B5'].value).strip().upper() == 'RAZON SOCIAL'), None)
    return proceso, vendedor


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('archivo')
    ap.add_argument('--maestro', required=True)
    ap.add_argument('--ids', nargs='*', default=[],
                    help='CSV ISBN;ID;TITULO con los titulos esperados, en orden')
    args = ap.parse_args()

    problemas, avisos = [], []
    wb = load_workbook(args.archivo)
    wsp, wsv = buscar_hojas(wb)
    print(f'hojas: {wb.sheetnames} | proceso: {wsp.title if wsp else "-"} '
          f'| planilla: {wsv.title if wsv else "-"} | primera: {wb.worksheets[0].title}')
    if wsp is None or wsv is None:
        raise SystemExit('*** no encuentro la hoja de proceso o la planilla del vendedor')
    if wb.worksheets[0] is not wsp:
        problemas.append(f'la hoja de proceso ({wsp.title}) no es la primera del libro')

    # --- 1. IDs ---
    ids_p = []
    c = 9
    while wsp.cell(1, c).value is not None:
        ids_p.append(wsp.cell(1, c).value)
        c += 1
    ids_v = [wsv.cell(5, 10 + i).value for i in range(len(ids_p) + 5)]
    ids_v = ids_v[:len(ids_v) - next((i for i, v in enumerate(reversed(ids_v)) if v is not None), 0)]
    if ids_p != ids_v:
        problemas.append(f'IDs distintos entre hojas: proceso {len(ids_p)}, planilla {len(ids_v)}')
    rep = {i for i in ids_p if ids_p.count(i) > 1}
    if rep:
        problemas.append(f'IDs repetidos: {sorted(rep)}')
    if args.ids:
        esperados = pd.concat([pd.read_csv(f, sep=';', dtype=str, encoding='utf-8-sig')
                               for f in args.ids])
        esp_ids = [int(x) for x in esperados['ID']]
        if ids_p != esp_ids:
            faltan = [i for i in esp_ids if i not in ids_p]
            sobran = [i for i in ids_p if i not in esp_ids]
            orden = not faltan and not sobran
            problemas.append('IDs de la fila de encabezado distintos a los del mes'
                             + (' (mismo conjunto, otro orden)' if orden else '')
                             + (f' | faltan {faltan}' if faltan else '')
                             + (f' | sobran {sobran}' if sobran else ''))
        isbn_de = dict(zip(esp_ids, esperados['ISBN']))
        for k, i in enumerate(ids_v):
            isbn = wsv.cell(3, 10 + k).value
            if i in isbn_de and str(isbn) != isbn_de[i]:
                problemas.append(f'{wsv.cell(3, 10 + k).coordinate}: ISBN {isbn} no corresponde '
                                 f'al ID {i} ({isbn_de[i]})')
    print(f'IDs: {len(ids_p)}')

    # --- clientes ---
    filas_p, r = [], 2
    while wsp.cell(r, 1).value is not None:
        filas_p.append(r)
        r += 1
    filas_v, r = [], 6
    while wsv.cell(r, 1).value is not None:
        filas_v.append(r)
        r += 1
    print(f'clientes: proceso {len(filas_p)}, planilla {len(filas_v)}')
    if len(filas_p) != len(filas_v):
        problemas.append(f'cantidad de clientes distinta: proceso {len(filas_p)}, planilla {len(filas_v)}')

    # --- 2. proceso == planilla sin la columna B ---
    n = len(ids_p)
    cols_v = [1] + list(range(3, 9 + n + 1))   # la planilla sin la columna B
    for rp, rv in zip(filas_p, filas_v):
        a = [wsp.cell(rp, c).value for c in range(1, 8 + n + 1)]
        b = [wsv.cell(rv, c).value for c in cols_v]
        for k, (x, y) in enumerate(zip(a, b)):
            if x != y:
                problemas.append(f'{wsp.cell(rp, k + 1).coordinate} ({x!r}) no coincide con la '
                                 f'planilla {wsv.cell(rv, cols_v[k]).coordinate} ({y!r})')

    # --- 3. cantidades y 4. columnas fijas ---
    listas = set()
    for rp in filas_p:
        cli = wsp.cell(rp, 1).value
        for k in range(n):
            v = wsp.cell(rp, 9 + k).value
            if not isinstance(v, (int, float)) or isinstance(v, bool):
                problemas.append(f'{wsp.cell(rp, 9 + k).coordinate} cliente {cli}, ID {ids_p[k]}: '
                                 f'cantidad {v!r} (tiene que ser un numero; si no pide, 0)')
            elif v < 0 or v != int(v):
                problemas.append(f'{wsp.cell(rp, 9 + k).coordinate} cliente {cli}: cantidad {v}')
        fijos = {3: 'Q', 4: 2, 5: 2}
        for col, val in fijos.items():
            if wsp.cell(rp, col).value != val:
                problemas.append(f'{wsp.cell(rp, col).coordinate} cliente {cli}: '
                                 f'{wsp.cell(1, col).value} = {wsp.cell(rp, col).value!r}, se espera {val!r}')
        if not isinstance(wsp.cell(rp, 1).value, int) or not isinstance(wsp.cell(rp, 2).value, int):
            problemas.append(f'fila {rp}: IdCliente o DirEnt no son numeros')
        listas.add(wsp.cell(rp, 7).value)
    if len(listas) != 1:
        problemas.append(f'mas de una lista de precios: {listas}')
    print(f'lista de precios: {listas}')

    # --- 5. maestro ---
    m = pd.read_excel(args.maestro)
    vendedor = str(wsv['A1'].value or '').strip()
    vistos = {}
    for rp, rv in zip(filas_p, filas_v):
        cli, dirent = wsp.cell(rp, 1).value, wsp.cell(rp, 2).value
        razon = wsv.cell(rv, 2).value
        vistos.setdefault(cli, []).append(rp)
        dm = m[m.IDCLIENTE == cli]
        if dm.empty:
            problemas.append(f'fila {rp}: cliente {cli} ({razon}) NO esta en el maestro')
            continue
        nombre = dm.iloc[0]['NOMBRE']
        if dm['FECHABAJA'].notna().all():
            problemas.append(f'fila {rp}: cliente {cli} ({nombre}) dado de baja el {dm.iloc[0]["FECHABAJA"]}')
        d = dm[dm.NROSECUENCIA == dirent]
        if d.empty:
            problemas.append(f'fila {rp}: cliente {cli} no tiene direccion {dirent} '
                             f'(tiene {sorted(dm.NROSECUENCIA.tolist())})')
        elif d.iloc[0]['TIPO DIRECCION'] != 'ENTR':
            entr = dm[dm['TIPO DIRECCION'] == 'ENTR'].NROSECUENCIA.tolist()
            problemas.append(f'fila {rp}: cliente {cli} DirEnt {dirent} es {d.iloc[0]["TIPO DIRECCION"]}, '
                             f'no ENTR (las de entrega son {entr})')
        if parecido(razon, nombre) < 0.6 and parecido(razon, dm.iloc[0]['FANTASIA']) < 0.6:
            problemas.append(f'fila {rp}: cliente {cli} en la planilla es "{razon}" y en el '
                             f'maestro "{nombre}"')
        vend_m = str(dm.iloc[0]['VENDEDOR'])
        if vendedor and parecido(vendedor, vend_m) < 0.6:
            avisos.append(f'fila {rp}: cliente {cli} ({nombre}) en el maestro es de "{vend_m}"')
    for cli, rs in vistos.items():
        if len(rs) > 1:
            dirs = [wsp.cell(r, 2).value for r in rs]
            (problemas if len(set(dirs)) < len(dirs) else avisos).append(
                f'cliente {cli} aparece {len(rs)} veces (filas {rs}, DirEnt {dirs})')

    # --- 6. cuotas ---
    ult = filas_v[-1] if filas_v else 5
    fila_cuota = next((r for r in range(ult + 1, ult + 6)
                       if any(isinstance(wsv.cell(r, 10 + k).value, (int, float)) for k in range(n))), None)
    if fila_cuota is None:
        avisos.append('no encuentro la fila de cuotas')
    elif fila_cuota != ult + 2:
        avisos.append(f'cuotas en la fila {fila_cuota}, ultimo cliente en la {ult} (regla: +2)')
    else:
        print(f'cuotas en la fila {fila_cuota} (ultimo cliente {ult} + 2)')

    print(f'\n{len(problemas)} PROBLEMAS')
    for p in problemas:
        print('  *** ' + p)
    print(f'\n{len(avisos)} avisos')
    for a in avisos:
        print('  -   ' + a)


if __name__ == '__main__':
    main()
