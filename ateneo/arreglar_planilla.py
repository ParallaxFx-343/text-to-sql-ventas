#!/usr/bin/env python3
"""Pasa a .xlsx las planillas de liquidacion que llegan "rotas".

Casos reales (octubre 2026):
- `liq_cons....xls` de clientes: BIFF2 (Excel 2.1, 1987) suelto, sin el
  contenedor OLE2. No esta danado, pero el Excel actual lo bloquea. Se lee
  con xlrd.
- `Liquidacion-....ods`: hoja con salto de linea en el nombre y mas de 31
  caracteres (Excel no la abre) y totales con formulas `oooc:=SUM([.I2:.I20])`
  sin valor guardado. Se lee el content.xml a mano (odfpy no instala aca) y
  las formulas se reescriben como =SUM(I2:I20).

No cambia ningun dato: copia los valores tal cual (los ISBN que vienen como
texto quedan como texto) y al final imprime un cuadre releyendo el xlsx.

Uso: python -P arreglar_planilla.py ENTRADA.(xls|ods) SALIDA.xlsx
"""
import datetime as dt
import re
import sys
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

NS = {'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
      'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
      'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}


def q(p, t):
    return '{%s}%s' % (NS[p], t)


def leer_ods(ruta):
    root = ET.fromstring(zipfile.ZipFile(ruta).read('content.xml'))
    tabla = next(root.iter(q('table', 'table')))
    nombre = tabla.get(q('table', 'name'))
    filas = []
    for row in tabla.iter(q('table', 'table-row')):
        f = []
        for c in row:
            if c.tag not in (q('table', 'table-cell'), q('table', 'covered-table-cell')):
                continue
            n = min(int(c.get(q('table', 'number-columns-repeated'), '1')), 50)
            vt = c.get(q('office', 'value-type'))
            formula = c.get(q('table', 'formula'))
            if formula:
                m = re.match(r'(?:oooc|of):=SUM\(\[\.([A-Z]+\d+):\.([A-Z]+\d+)\]\)', formula)
                v = f'=SUM({m.group(1)}:{m.group(2)})' if m else None
                if v is None:
                    print(f'  *** formula que no se tradujo: {formula}')
            elif vt in ('float', 'currency', 'percentage'):
                v = float(c.get(q('office', 'value')))
            elif vt == 'date':
                v = dt.datetime.strptime(c.get(q('office', 'date-value'))[:10], '%Y-%m-%d')
            elif vt is None:
                v = None
            else:
                v = ' '.join(''.join(p.itertext()) for p in c.findall(q('text', 'p'))) or None
            f += [v] * n
        while f and f[-1] is None:
            f.pop()
        filas.append(f)
    while filas and not filas[-1]:
        filas.pop()
    return nombre, filas


def leer_xls(ruta):
    import xlrd
    sh = xlrd.open_workbook(ruta, encoding_override='cp1252').sheet_by_index(0)
    filas = []
    for r in range(sh.nrows):
        f = []
        for c in range(sh.ncols):
            t, v = sh.cell_type(r, c), sh.cell_value(r, c)
            f.append(None if v == '' else v)
        filas.append(f)
    return sh.name, filas


def main():
    entrada, salida = sys.argv[1], sys.argv[2]
    nombre, filas = (leer_ods if entrada.lower().endswith('.ods') else leer_xls)(entrada)
    limpio = re.sub(r'\s*\(.*$', '', re.sub(r'\s+', ' ', nombre or 'Hoja1')).strip()[:31] or 'Hoja1'
    wb = Workbook()
    ws = wb.active
    ws.title = re.sub(r'[\[\]:*?/\\]', '', limpio)
    for i, f in enumerate(filas, 1):
        for j, v in enumerate(f, 1):
            if isinstance(v, float) and v == int(v):
                v = int(v)
            c = ws.cell(i, j, v)
            if isinstance(v, dt.datetime):
                c.number_format = 'DD-MM-YYYY'
            if i == 1 or (isinstance(v, str) and v.startswith('=')):
                c.font = Font(bold=True)
    ws.freeze_panes = 'A2'
    wb.save(salida)
    # cuadre releyendo
    rel = load_workbook(salida).active
    iguales = all(rel.cell(i, j).value == (int(v) if isinstance(v, float) and v == int(v) else v)
                  for i, f in enumerate(filas, 1) for j, v in enumerate(f, 1))
    print(f'hoja "{nombre!r}" -> "{ws.title}" | {len(filas)} filas | xlsx == origen: {iguales}')


if __name__ == '__main__':
    main()
