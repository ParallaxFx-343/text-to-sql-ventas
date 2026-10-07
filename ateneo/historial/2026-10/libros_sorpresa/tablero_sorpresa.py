"""Tablero de ventas de los Libros Sorpresa por local (solo locales fisicos).

Uso: python -P tablero_sorpresa.py DETALLE.xlsx AGRUPADAS.xlsx config.json SALIDA.xlsx [SII.xlsx]
- DETALLE: "Detalle Venta Articulo" (una fila por venta: Sucursal "NOMBRE  (nro)",
  Fecha, Articulo, Titulo, Cantidad, Bruto, Total, ISBN)
- AGRUPADAS: "Ventas Proveedor Agrupadas Articulo" (solo para cuadrar y para
  saber que locales vendieron Editorial pero no Sorpresa)
- config.json: padron de locales fisicos (orden_paec)
- SII (opcional): stock por local, para "cuanto le queda" y "le alcanza" (dias).
  Del deposito no se muestra nada (Franco: el stock de deposito se va a liberar).
"""
import json
import sys
from datetime import date

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

f_det, f_agr, f_cfg, salida = sys.argv[1:5]
f_sii = sys.argv[5] if len(sys.argv) > 5 else None
padron = set(json.load(open(f_cfg, encoding='utf-8'))['orden_paec']) | {'DEVOTO'}

# ---------- datos ----------
det = pd.read_excel(f_det, dtype={'ISBN': str})
det['nro'] = det.Sucursal.str.extract(r'\((\d+)\)\s*$')[0].astype(int)
det['local'] = det.Sucursal.str.replace(r'\s*\(\d+\)\s*$', '', regex=True).str.strip()
det = det[det.local.isin(padron)].copy()          # solo locales fisicos
NOMBRES = {'HISTORICA': 'Histórica', 'ROMANTICA': 'Romántica', 'CONTEMPORANEA': 'Contemporánea',
           'MISTERIO': 'Misterio', 'EROTICA': 'Erótica', 'BIO': 'Bío'}
det['titulo'] = det.Titulo.str.replace('LIBRO SORPRESA ', '', regex=False).str.strip().map(NOMBRES)
assert det.titulo.notna().all(), 'hay un Libro Sorpresa nuevo sin nombre en NOMBRES'
det['dia'] = det.Fecha.dt.date

agr = pd.read_excel(f_agr, dtype={'Código / ISBN': str})
agr = agr[agr.Canal.isna() | (agr.Canal.astype(str).str.strip() == '')]   # sin canales online
agr = agr[agr['Nombre Sucursal'].isin(padron)]
sor = agr[agr.TITULO.astype(str).str.contains('LIBRO SORPRESA')]

# cuadre por dos caminos: detalle vs agrupadas, local x articulo
a = sor.groupby(['Código Sucursal', 'ID ARTICULO']).Cantidad.sum()
d = det.groupby(['nro', 'Articulo']).Cantidad.sum()
a.index.names = d.index.names = ['s', 'a']
cuadre = pd.concat([a, d], axis=1).fillna(0)
assert (cuadre.iloc[:, 0] == cuadre.iloc[:, 1]).all(), 'detalle y agrupadas no cuadran'

dias = sorted(det.dia.unique())
ult = dias[-1]                                      # dia parcial (export a la mañana)
hasta = max(det.Fecha)                              # para el subtitulo
DIAS_ES = ['Lun', 'Mar', 'Mié', 'Jue', 'Vie', 'Sáb', 'Dom']
etq = {x: f'{DIAS_ES[x.weekday()]} {x:%d/%m}' for x in dias}

tot_u = int(det.Cantidad.sum())
tot_ars = int(det.Total.sum())
por_dia = det.groupby('dia').Cantidad.sum().reindex(dias, fill_value=0)
por_tit = det.groupby('titulo').Cantidad.sum().sort_values(ascending=False)
mat = det.pivot_table(index=['nro', 'local'], columns='dia', values='Cantidad',
                      aggfunc='sum', fill_value=0).reindex(columns=dias, fill_value=0)
mat['Total'] = mat.sum(axis=1)
primera = det[det.Cantidad > 0].groupby(['nro', 'local']).dia.min()
ult_completo = dias[-2]
def ritmo(row):
    p = primera[row.name]
    n = (ult_completo - p).days + 1
    u = sum(row[x] for x in dias if p <= x <= ult_completo)
    return round(u / n, 1) if n > 0 else None
mat['Primera venta'] = [primera[i] for i in mat.index]
mat['Ritmo'] = mat.apply(ritmo, axis=1)
devol = det[det.Cantidad < 0]
mat['Devol.'] = devol.groupby(['nro', 'local']).Cantidad.sum().reindex(mat.index, fill_value=0)
mat = mat.sort_values(['Total', 'Ritmo'], ascending=False)
tit_cols = list(por_tit.index)
mt = det.pivot_table(index=['nro', 'local'], columns='titulo', values='Cantidad',
                     aggfunc='sum', fill_value=0).reindex(columns=tit_cols, fill_value=0)
mt['Total'] = mt.sum(axis=1)
mt = mt.loc[mat.index]

# ---------- stock en locales (SII) ----------
stock = None
if f_sii:
    sii = pd.read_excel(f_sii)
    cols_s = list(sii.columns)
    suc_s = [c for c in cols_s[cols_s.index('CLASIFIC') + 1:] if c in padron]
    assert len(suc_s) == len(padron), f'faltan locales en el SII: {padron - set(suc_s)}'
    s = sii[sii.TITULO.astype(str).str.startswith('LIBRO SORPRESA')].copy()
    s['titulo'] = s.TITULO.str.replace('LIBRO SORPRESA ', '', regex=False).str.strip().map(NOMBRES)
    assert s.titulo.notna().all() and s.titulo.is_unique and len(s) == len(NOMBRES)
    stock = s.set_index('titulo')[suc_s].fillna(0).astype(int).T   # local x titulo
    stock = stock.reindex(columns=list(por_tit.index))
    stock['Total'] = stock.sum(axis=1)
    loc_de = {loc: (n, loc) for n, loc in mat.index}
    mat['Stock'] = [int(stock.loc[loc, 'Total']) for _, loc in mat.index]
    mat['Alcanza'] = [round(r.Stock / r.Ritmo, 1) if r.Ritmo else None for _, r in mat.iterrows()]
    import re as _re
    m_fecha = _re.search(r'(20\d{2})(\d{2})(\d{2})(\d{2})(\d{2})', f_sii.split('/')[-1])
    fecha_sii = f'{m_fecha.group(3)}/{m_fecha.group(2)} {m_fecha.group(4)}:{m_fecha.group(5)}' if m_fecha else 'SII'
    llenos = [x for x in sorted(det.dia.unique())[:-1]]          # sin el dia parcial
    vend_d = det[det.dia.isin(llenos)].groupby('titulo').Cantidad.sum() / len(llenos)
    alcanza_tit = {t: (stock[t].sum() / vend_d[t]) for t in por_tit.index if vend_d.get(t, 0) > 0}

con_venta = set(det.local)
sin_venta = sorted(padron - con_venta)
otra_editorial = set(agr['Nombre Sucursal'])       # vendieron Editorial en el periodo
n_locales = len(con_venta)
n_padron = len(padron)

assert int(mat.Total.sum()) == tot_u == int(mt.Total.sum()) == int(por_tit.sum()) == int(por_dia.sum())

# ---------- estilo ----------
AZUL, AZUL_OSC, GRIS = '2A78D6', '1C5CAB', 'C3C2B7'
TINTA, TINTA2, FONDO, BORDE = '0B0B0B', '52514E', 'F5F6F8', 'D9DCE1'
fino = Side(style='thin', color=BORDE)
caja = Border(left=fino, right=fino, top=fino, bottom=fino)
H = Font(name='Calibri', bold=True, color='FFFFFF', size=10)
HF = PatternFill('solid', fgColor='1F4E78')
def cab(ws, fila, textos, col0=1):
    for i, t in enumerate(textos):
        c = ws.cell(fila, col0 + i, t)
        c.font, c.fill, c.border = H, HF, caja
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
def calor(ws, rango):
    ws.conditional_formatting.add(rango, ColorScaleRule(
        start_type='num', start_value=0, start_color='FFFFFF',
        end_type='max', end_color='6DA7EC'))

wb = Workbook()

# ---------- hoja Resumenes (fuente de los graficos y vista de tabla) ----------
rs = wb.active
rs.title = 'Resúmenes'
rs['A1'] = 'Unidades por día'
rs['A1'].font = Font(bold=True, size=12)
cab(rs, 2, ['Día', 'Unidades'])
for i, x in enumerate(dias, start=3):
    rs.cell(i, 1, etq[x] + (' (parcial)' if x == ult else ''))
    rs.cell(i, 2, int(por_dia[x]))
fd = 2 + len(dias)
rs['D1'] = 'Unidades por título'
rs['D1'].font = Font(bold=True, size=12)
cab(rs, 2, ['Título', 'Unidades', '% del total', 'Locales'], col0=4)
for i, (t, u) in enumerate(por_tit.items(), start=3):
    rs.cell(i, 4, t); rs.cell(i, 5, int(u))
    rs.cell(i, 6, u / tot_u).number_format = '0%'
    rs.cell(i, 7, int(det[(det.titulo == t) & (det.Cantidad > 0)].local.nunique()))
ft = 2 + len(por_tit)
rs['I1'] = 'Top 15 locales'
rs['I1'].font = Font(bold=True, size=12)
cab(rs, 2, ['Local', 'Unidades'], col0=9)
top = mat.Total.head(15)[::-1]                      # al reves: la barra mas larga arriba
for i, ((n, loc), u) in enumerate(top.items(), start=3):
    rs.cell(i, 9, loc.title()); rs.cell(i, 10, int(u))
fl = 2 + len(top)
for c, w in zip('ABCDEFGHIJ', (16, 10, 3, 16, 10, 11, 9, 3, 30, 10)):
    rs.column_dimensions[c].width = w

# ---------- hoja Tablero ----------
tb = wb.create_sheet('Tablero', 0)
tb.sheet_view.showGridLines = False
for c in range(1, 17):
    tb.column_dimensions[L(c)].width = 10.5
for r in range(1, 70):
    for c in range(1, 17):
        tb.cell(r, c).fill = PatternFill('solid', fgColor=FONDO)
tb.merge_cells('A1:P1')
tb['A1'] = 'Libros Sorpresa · ventas por local'
tb['A1'].font = Font(name='Calibri', bold=True, size=20, color=TINTA)
tb.merge_cells('A2:P2')
tb['A2'] = (f'Solo locales físicos · del {dias[0]:%d/%m} al {ult:%d/%m/%Y} · '
            f'el {ult:%d/%m} es parcial (el reporte se sacó ese mismo día a la mañana) · '
            f'fuente: Detalle Venta Artículo, cuadrado contra Ventas Proveedor Agrupadas')
tb['A2'].font = Font(size=10, color=TINTA2)
tb.row_dimensions[1].height = 30

mejor = por_dia[por_dia.index != ult].idxmax()
kpis = [
    ('Unidades netas', f'{tot_u:,}'.replace(',', '.'), f'{len(devol)} devoluciones ya restadas'),
    ('Facturación', '$ ' + f'{tot_ars:,}'.replace(',', '.'), f'a $ {int(det.Total.sum() / tot_u):,} c/u'.replace(',', '.')),
    ('Locales con venta', f'{n_locales} de {n_padron}', f'{len(sin_venta)} sin venta (ver hoja)'),
    ('Mejor día', f'{etq[mejor]}', f'{int(por_dia[mejor])} u · {por_dia[mejor] / tot_u:.0%} del total'),
    (('Le queda en los locales', f'{int(stock.Total.sum()):,}'.replace(',', '.') + ' u',
      f'SII {fecha_sii} · ver hoja Le quedan x título') if stock is not None else
     ('Promedio por local', f'{tot_u / n_locales:.1f}'.replace('.', ','), 'unidades en el período')),
]
for k, (lab, val, nota) in enumerate(kpis):
    c0 = 1 + k * 3 + (k > 0) * 0
    c0 = [1, 4, 7, 10, 13][k]
    for r in (4, 5, 6):
        tb.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c0 + 2)
        for c in range(c0, c0 + 3):
            tb.cell(r, c).fill = PatternFill('solid', fgColor='FFFFFF')
            tb.cell(r, c).border = Border(left=fino if c == c0 else None, right=fino if c == c0 + 2 else None,
                                          top=fino if r == 4 else None, bottom=fino if r == 6 else None)
    tb.cell(4, c0, lab).font = Font(size=10, color=TINTA2)
    tb.cell(5, c0, val).font = Font(size=20, bold=True, color=TINTA)
    tb.cell(6, c0, nota).font = Font(size=9, color=TINTA2)
    for r in (4, 5, 6):
        tb.cell(r, c0).alignment = Alignment(horizontal='left', vertical='center', indent=1)
tb.row_dimensions[5].height = 30

def barras(titulo, cats, vals, horizontal=False, ancho=15, alto=7.5):
    ch = BarChart()
    ch.type = 'bar' if horizontal else 'col'
    ch.title = titulo
    ch.style = 10
    ch.legend = None
    ch.add_data(vals, titles_from_data=False)
    ch.set_categories(cats)
    s = ch.series[0]
    s.graphicalProperties = GraphicalProperties(solidFill=AZUL)
    s.graphicalProperties.line = LineProperties(noFill=True)
    ch.gapWidth = 60
    ch.dataLabels = DataLabelList(showVal=True, showCatName=False, showSerName=False,
                                   showLegendKey=False, showPercent=False, showLeaderLines=False)
    ch.y_axis.majorGridlines = None
    ch.y_axis.delete = True
    ch.x_axis.delete = False
    ch.width, ch.height = ancho, alto
    return ch

c1 = barras('Unidades por día', Reference(rs, min_col=1, min_row=3, max_row=fd),
            Reference(rs, min_col=2, min_row=3, max_row=fd))
p = DataPoint(idx=len(dias) - 1)                    # el dia parcial en gris
p.graphicalProperties = GraphicalProperties(solidFill=GRIS)
c1.series[0].dPt.append(p)
tb.add_chart(c1, 'A8')
c2 = barras('Unidades por título', Reference(rs, min_col=4, min_row=3, max_row=ft),
            Reference(rs, min_col=5, min_row=3, max_row=ft))
tb.add_chart(c2, 'I8')
c3 = barras('Top 15 locales (unidades)', Reference(rs, min_col=9, min_row=3, max_row=fl),
            Reference(rs, min_col=10, min_row=3, max_row=fl), horizontal=True, ancho=31, alto=10)
tb.add_chart(c3, 'A24')

# lectura
r0 = 45
tb.cell(r0, 1, 'Lo que muestran los datos').font = Font(bold=True, size=13, color=TINTA)
arranque = primera.value_counts().sort_index()
lideres = por_tit.head(2)
bio = por_tit.get('Bío', 0)
lineas = [
    f'{lideres.index[0]} ({int(lideres.iloc[0])} u, {lideres.iloc[0] / tot_u:.0%}) y {lideres.index[1]} '
    f'({int(lideres.iloc[1])} u) son los más vendidos; Bío quedó muy abajo ({int(bio)} u, vendido en '
    f'{det[(det.titulo == "Bío") & (det.Cantidad > 0)].local.nunique()} de {n_locales} locales).',
    f'El pico fue el {etq[mejor]} con {int(por_dia[mejor])} u ({por_dia[mejor] / tot_u:.0%} de todo lo vendido); '
    'después: ' + ', '.join(f'{etq[x]} {int(por_dia[x])} u' for x in dias if mejor < x < ult) +
    '. Con tan pocos días todavía no se puede hablar de tendencia.',
    'Los locales no arrancaron todos juntos: ' + ', '.join(
        f'{int(n)} el {x:%d/%m}' for x, n in arranque.items()) +
    '. Para comparar locales que arrancaron en días distintos, usar la columna "Vendió por día" de la hoja Local x día.',
]
if stock is None:
    lineas.append(
        f'Sin ninguna venta de Sorpresa: {", ".join(s.title() for s in sin_venta)}. Todos menos Devoto sí '
        'vendieron otros títulos de la Editorial según Ventas Agrupadas, así que conviene confirmar si les llegaron. '
        'A Devoto la repo se le pasó recién el 05/10.')
else:
    cero = [s for s in sin_venta if stock.loc[s, 'Total'] == 0]
    con = [s for s in sin_venta if stock.loc[s, 'Total'] > 0]
    orden = sorted(alcanza_tit.items(), key=lambda x: x[1])
    lineas.append(
        f'Lo que le queda a la cadena (SII {fecha_sii}), al ritmo promedio de la primera semana, alcanza para: '
        + ', '.join(f'{t} {d:.1f} días'.replace('.', ',') for t, d in orden[:3])
        + f'; {orden[-1][0]}, en cambio, tiene para {orden[-1][1]:.0f} días. El sábado se vendió '
        + f'{por_dia[mejor] / por_dia[[x for x in dias if x != ult]].mean():.1f}'.replace('.', ',')
        + ' veces el promedio diario, así que un fin de semana consume bastante más que eso.')
    lineas.append(
        'Sin ninguna venta: ' + ', '.join(s.title() for s in cero) + ' no tienen stock de Sorpresa (no les llegó); '
        + ', '.join(f'{s.title()} ({int(stock.loc[s, "Total"])} u)' for s in con)
        + ' sí tienen stock y no vendieron' + (' (a Devoto la repo se le pasó el 05/10)' if 'DEVOTO' in con else '') + '.')
for i, t in enumerate(lineas):
    tb.merge_cells(start_row=r0 + 1 + i, start_column=1, end_row=r0 + 1 + i, end_column=16)
    c = tb.cell(r0 + 1 + i, 1, '•  ' + t)
    c.font = Font(size=10.5, color=TINTA)
    c.alignment = Alignment(wrap_text=True, vertical='top')
    tb.row_dimensions[r0 + 1 + i].height = 30

# ---------- hoja Local x dia ----------
ld = wb.create_sheet('Local x día', 1)
ld.freeze_panes = 'C3'
ld['A1'] = (f'Unidades por local y por día · ordenado por total · el {ult:%d/%m} es parcial · '
            f'"Vendió por día" = unidades por día desde su primera venta hasta el {ult_completo:%d/%m}'
            + (f' · "Le quedan" = stock del local en el SII del {fecha_sii} · "Le alcanza" = le quedan / vendió por día; '
               'en rojo, 3 días o menos' if stock is not None else ''))
ld['A1'].font = Font(italic=True, size=9, color=TINTA2)
cols = ['Nro', 'Local'] + [etq[x] + ('*' if x == ult else '') for x in dias] + \
       ['Total', 'Primera venta', 'Vendió por día'] + \
       (['Le quedan', 'Le alcanza (días)'] if stock is not None else []) + ['Devol.']
extra = 2 if stock is not None else 0
cab(ld, 2, cols)
for i, ((n, loc), row) in enumerate(mat.iterrows(), start=3):
    vals = [n, loc] + [int(row[x]) for x in dias] + [int(row.Total), row['Primera venta'], row.Ritmo] + \
           ([int(row.Stock), row.Alcanza] if stock is not None else []) + [int(row['Devol.'])]
    for j, v in enumerate(vals, start=1):
        c = ld.cell(i, j, v); c.border = caja
        c.alignment = Alignment(horizontal='left' if j == 2 else 'center')
    ld.cell(i, 3 + len(dias)).font = Font(bold=True)
    ld.cell(i, 4 + len(dias)).number_format = 'DD/MM'
    ld.cell(i, 5 + len(dias)).number_format = '0.0'
    if stock is not None:
        ld.cell(i, 7 + len(dias)).number_format = '0.0'
    if row['Devol.'] == 0:
        ld.cell(i, 6 + len(dias) + extra).value = None
uf = 2 + len(mat)
tr = uf + 1
ld.cell(tr, 2, 'TOTAL').font = Font(bold=True)
for j in list(range(3, 4 + len(dias))) + ([6 + len(dias)] if stock is not None else []):
    c = ld.cell(tr, j, f'=SUM({L(j)}3:{L(j)}{uf})')
    c.font = Font(bold=True); c.border = caja; c.alignment = Alignment(horizontal='center')
calor(ld, f'C3:{L(2 + len(dias))}{uf}')
if stock is not None:
    from openpyxl.formatting.rule import CellIsRule
    col_alc = L(7 + len(dias))
    ld.conditional_formatting.add(f'{col_alc}3:{col_alc}{uf}', CellIsRule(
        operator='lessThanOrEqual', formula=['3'], fill=PatternFill('solid', fgColor='F8D7D7'),
        font=Font(bold=True, color='9C1C1C')))
for c, w in zip(range(1, len(cols) + 1), [6, 30] + [10] * len(dias) + [8, 10, 9] + [9, 10] * (extra // 2) + [7]):
    ld.column_dimensions[L(c)].width = w

# ---------- hoja Local x titulo ----------
lt = wb.create_sheet('Local x título', 2)
lt.freeze_panes = 'C3'
lt['A1'] = 'Unidades por local y por título · mismo orden que Local x día'
lt['A1'].font = Font(italic=True, size=9, color=TINTA2)
cab(lt, 2, ['Nro', 'Local'] + tit_cols + ['Total'])
for i, ((n, loc), row) in enumerate(mt.iterrows(), start=3):
    for j, v in enumerate([n, loc] + [int(row[t]) for t in tit_cols] + [int(row.Total)], start=1):
        c = lt.cell(i, j, v); c.border = caja
        c.alignment = Alignment(horizontal='left' if j == 2 else 'center')
    lt.cell(i, 3 + len(tit_cols)).font = Font(bold=True)
uf2 = 2 + len(mt)
lt.cell(uf2 + 1, 2, 'TOTAL').font = Font(bold=True)
for j in range(3, 4 + len(tit_cols)):
    c = lt.cell(uf2 + 1, j, f'=SUM({L(j)}3:{L(j)}{uf2})')
    c.font = Font(bold=True); c.border = caja; c.alignment = Alignment(horizontal='center')
calor(lt, f'C3:{L(2 + len(tit_cols))}{uf2}')
for c, w in zip(range(1, len(tit_cols) + 4), [6, 30] + [13] * len(tit_cols) + [8]):
    lt.column_dimensions[L(c)].width = w

# ---------- hoja Sin venta ----------
sv = wb.create_sheet('Sin venta', 3)
cab(sv, 1, ['Local físico', '¿Tiene ventas de otros títulos de la Editorial en Ventas Agrupadas?']
    + ([f'Le quedan Sorpresa ({fecha_sii})'] if stock is not None else []))
for i, s in enumerate(sin_venta, start=2):
    sv.cell(i, 1, s).border = caja
    txt = 'Sí' if s in otra_editorial else 'No'
    if s == 'DEVOTO':
        txt += ' (la repo de Sorpresa se pasó el 05/10)'
    sv.cell(i, 2, txt).border = caja
    if stock is not None:
        c = sv.cell(i, 3, int(stock.loc[s, 'Total'])); c.border = caja
        c.alignment = Alignment(horizontal='center')
sv.column_dimensions['A'].width = 30
sv.column_dimensions['B'].width = 55
sv.column_dimensions['C'].width = 22

# ---------- hoja Stock x titulo ----------
if stock is not None:
    sk = wb.create_sheet('Le quedan x título', 3)
    sk.freeze_panes = 'C3'
    sk['A1'] = (f'Stock de cada local en el SII del {fecha_sii}, por título · mismo orden que Local x día; '
                'al final los locales sin venta · abajo, el total de la cadena y para cuántos días alcanza')
    sk['A1'].font = Font(italic=True, size=9, color=TINTA2)
    cab(sk, 2, ['Nro', 'Local'] + tit_cols + ['Total'])
    filas = [(n, loc) for n, loc in mat.index] + [('', s) for s in sin_venta]
    for i, (n, loc) in enumerate(filas, start=3):
        for j, v in enumerate([n, loc] + [int(stock.loc[loc, t]) for t in tit_cols]
                              + [int(stock.loc[loc, 'Total'])], start=1):
            c = sk.cell(i, j, v); c.border = caja
            c.alignment = Alignment(horizontal='left' if j == 2 else 'center')
        sk.cell(i, 3 + len(tit_cols)).font = Font(bold=True)
    uf3 = 2 + len(filas)
    calor(sk, f'C3:{L(2 + len(tit_cols))}{uf3}')
    dias_llenos = [x for x in dias if x != ult]
    vend = det[det.dia.isin(dias_llenos)].groupby('titulo').Cantidad.sum()
    resumen = [('Le queda en la cadena', [int(stock[t].sum()) for t in tit_cols] + [int(stock.Total.sum())], '#,##0'),
               (f'Vendido {dias_llenos[0]:%d/%m} al {dias_llenos[-1]:%d/%m}',
                [int(vend.get(t, 0)) for t in tit_cols] + [int(vend.sum())], '#,##0'),
               ('Vendido por día (promedio)', [round(vend.get(t, 0) / len(dias_llenos), 1) for t in tit_cols]
                + [round(vend.sum() / len(dias_llenos), 1)], '0.0'),
               ('Le alcanza (días)', [round(stock[t].sum() / (vend.get(t, 0) / len(dias_llenos)), 1)
                                      if vend.get(t, 0) else None for t in tit_cols]
                + [round(stock.Total.sum() / (vend.sum() / len(dias_llenos)), 1)], '0.0')]
    for k, (lab, vals, fmt) in enumerate(resumen):
        r = uf3 + 2 + k
        sk.cell(r, 2, lab).font = Font(bold=True)
        for j, v in enumerate(vals, start=3):
            c = sk.cell(r, j, v); c.font = Font(bold=True); c.border = caja
            c.number_format = fmt; c.alignment = Alignment(horizontal='center')
    from openpyxl.formatting.rule import CellIsRule
    sk.conditional_formatting.add(f'C{uf3 + 5}:{L(3 + len(tit_cols))}{uf3 + 5}', CellIsRule(
        operator='lessThanOrEqual', formula=['3'], fill=PatternFill('solid', fgColor='F8D7D7'),
        font=Font(bold=True, color='9C1C1C')))
    for c, w in zip(range(1, len(tit_cols) + 4), [6, 30] + [13] * len(tit_cols) + [8]):
        sk.column_dimensions[L(c)].width = w
    sk.column_dimensions['B'].width = 30
    sk.page_setup.orientation = 'landscape'
    sk.sheet_properties.pageSetUpPr.fitToPage = True
    sk.page_setup.fitToWidth = 1
    sk.page_setup.fitToHeight = 0
    cob_tit = dict(zip(tit_cols, resumen[3][1][:-1]))
    cob_tot = resumen[3][1][-1]

# ---------- hoja Devoluciones ----------
dv = wb.create_sheet('Devoluciones', 5)
cab(dv, 1, ['Local', 'Fecha', 'Título', 'Cantidad', 'Total $'])
for i, r in enumerate(devol.sort_values(['Fecha', 'local']).itertuples(), start=2):
    for j, v in enumerate([r.local, r.Fecha, r.titulo, int(r.Cantidad), int(r.Total)], start=1):
        c = dv.cell(i, j, v); c.border = caja
    dv.cell(i, 2).number_format = 'DD/MM/YYYY'
for c, w in zip('ABCDE', (30, 12, 16, 10, 10)):
    dv.column_dimensions[c].width = w

# ---------- hoja Datos ----------
dt = wb.create_sheet('Datos')
cab(dt, 1, ['Nro', 'Local', 'Fecha', 'Artículo', 'ISBN', 'Título', 'Cantidad', 'Total $'])
for i, r in enumerate(det.sort_values(['local', 'Fecha']).itertuples(), start=2):
    for j, v in enumerate([r.nro, r.local, r.Fecha, r.Articulo, r.ISBN, r.titulo,
                           int(r.Cantidad), int(r.Total)], start=1):
        dt.cell(i, j, v)
    dt.cell(i, 3).number_format = 'DD/MM/YYYY'
dt.auto_filter.ref = f'A1:H{len(det) + 1}'
dt.freeze_panes = 'A2'
for c, w in zip('ABCDEFGH', (6, 30, 12, 10, 16, 16, 10, 10)):
    dt.column_dimensions[c].width = w

# impresion: el tablero en una hoja apaisada; las tablas al ancho de la pagina
for ws, alto in ((tb, 1), (ld, 0), (lt, 0)):
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = alto
tb.print_area = 'A1:P50'
for nombre in ('Local x día', 'Local x título', 'Le quedan x título'):
    if nombre in wb.sheetnames:
        wb[nombre].print_title_rows = '2:2'     # encabezado en cada pagina impresa
wb.save(salida)
print(f'{tot_u} u | $ {tot_ars:,} | {n_locales} locales de {n_padron} | {len(dias)} dias | '
      f'devol {len(devol)} | sin venta {sin_venta} | cuadre agrupadas OK ({len(cuadre)} celdas)')
print('por dia', {etq[x]: int(v) for x, v in por_dia.items()})
print('por titulo', por_tit.to_dict())
for t in lineas: print(' -', t)
