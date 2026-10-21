"""Repo 21/10: un PAEC por sucursal. Reglas de la planilla: puestos 1-10 = cantidad exacta de Ventas;
puestos 11-40 = lo necesario para completar el stock a 10 (10 - Stock de la planilla, si da > 0).
Sin Libros Sorpresa (EAN 290) ni los 29 títulos de la lista "Libro sorpresa". Sin stock: van igual, marcados.
Uso: python -P repo2110_out.py BLOCKS.pkl SII.xls config.json SALIDA/ excluir.xlsx 'SII dd/mm hh:mm'"""
import sys, json, pickle, re, zipfile, pandas as pd
from pathlib import Path
from openpyxl import load_workbook
blocks = pickle.load(open(sys.argv[1], 'rb'))
orden = json.load(open(sys.argv[3]))['orden_paec']
out = Path(sys.argv[4]); out.mkdir(parents=True, exist_ok=True)
sii = pd.read_excel(sys.argv[2]); sii = sii[sii['ID ARTICULO'].notna()].copy(); sii['n'] = pd.to_numeric(sii.ISBN, errors='coerce'); by = sii.set_index('n')
EXSOR = {int(r[0]): r[6] for r in load_workbook(sys.argv[5])['Libro sorpresa'].iter_rows(min_row=5, values_only=True) if r[0]}
etq = sys.argv[6]
blocks.sort(key=lambda b: orden.index(b['suc']))
def cant(x):
    _, o, v, t, isbn, st = x
    return int(v) if o <= 10 else max(0, 10 - int(st or 0))
dem = {}
for b in blocks:
    for x in b['lineas']:
        i = int(x[4])
        if not str(i).startswith('290') and i not in EXSOR and cant(x) > 0: dem[i] = dem.get(i, 0) + cant(x)
parcial = {i: d for i, d in dem.items() if 0 < by.loc[i, 'DISPONIBLE'] < d}
T = ['REPOSICION 21/10 - UN PAEC POR SUCURSAL',
     'Puestos 1 a 10: la cantidad exacta de "Ventas". Puestos 11 a 40: lo necesario para completar el stock a 10 (10 - Stock de la planilla).',
     'Sin Libros Sorpresa ni los 29 títulos de la lista "Libro sorpresa". Los títulos sin disponible van igual, marcados [SIN STOCK].',
     'Con [DISP x PARA y EN TOTAL]: hay disponible pero no alcanza para lo que piden todas las sucursales juntas.',
     f'Disponible según {etq}. Orden de las sucursales: el de los TXT de PAEC.', '']
res = []; files = []
for b in blocks:
    sor = [x for x in b['lineas'] if str(x[4]).startswith('290')]
    lis = [x for x in b['lineas'] if int(x[4]) in EXSOR]
    L = [(x, cant(x)) for x in b['lineas'] if x not in sor and x not in lis]
    cero = [x for x, q in L if q == 0]
    L = [(x, q) for x, q in L if q > 0]
    slug = re.sub(r'[^A-Z0-9]+', '_', b['suc']).strip('_')
    p = out / f'reposicion_{slug}.csv'; files.append(p)
    p.write_bytes(('\r\n'.join(f'{x[4]};{q}' for x, q in L) + '\r\n').encode())
    ss = [(x, q) for x, q in L if by.loc[int(x[4]), 'DISPONIBLE'] <= 0]
    u = sum(q for _, q in L); u10 = sum(q for x, q in L if x[1] <= 10)
    T += ['=' * 72, f'{b["suc"]:<52}{u:>6} u  {len(L)} títulos  (puestos 1-10: {u10} u; 11-40: {u - u10} u)',
          f'   archivo {p.name} | sacados {len(sor)} Libros Sorpresa ({sum(cant(x) for x in sor)} u)'
          + (f' y {len(lis)} de la lista sorpresa (' + ', '.join(f'{x[3]} {cant(x)}' for x in lis) + ')' if lis else '')
          + f' | {len(cero)} puestos 11-40 ya tienen 10 o más | sin stock {len(ss)} ({sum(q for _, q in ss)} u)', '=' * 72]
    for x, q in L:
        d = int(by.loc[int(x[4]), 'DISPONIBLE'])
        reg = f'vendió {int(x[2])}' if x[1] <= 10 else f'stock {int(x[5] or 0)} -> 10'
        T.append(f'   {q:>4}  {x[4]}  {str(by.loc[int(x[4]), "TITULO"])[:40]:<40} #{x[1]:<2} {reg:<12}'
                 + ('  [SIN STOCK]' if d <= 0 else (f'  [DISP {d} PARA {parcial[int(x[4])]} EN TOTAL]' if int(x[4]) in parcial else '')))
    T.append('')
    res.append(dict(suc=b['suc'], lineas=len(L), u=u, u10=u10, sor=sum(cant(x) for x in sor), lis=[(x[3], cant(x)) for x in lis], ss=sum(q for _, q in ss), cero=len(cero)))
tl, tu, tss = sum(r['lineas'] for r in res), sum(r['u'] for r in res), sum(r['ss'] for r in res)
T += ['=' * 72, f'TOTAL {len(blocks)} sucursales | {tl} líneas | {tu} u (de esas, {tss} sin stock) | Sorpresa sacados: {sum(r["sor"] for r in res)} u']
(out / 'DETALLE_REPOSICION_21-10.txt').write_text('\n'.join(T) + '\n', encoding='utf-8-sig')
(out / 'PAECS_REPOSICION_21-10.txt').write_text('\n'.join(f'{b["suc"]}: ' for b in sorted(blocks, key=lambda b: b['suc'])) + '\n', encoding='utf-8-sig')
with zipfile.ZipFile(out / 'reposicion_21-10_por_sucursal.zip', 'w') as z:
    for p in files: z.write(p, p.name)
pd.DataFrame(res).to_json(out / 'resumen.json', orient='records', force_ascii=False)
for r in res: print(r)
print('TOTAL', tl, tu, 'sin stock', tss, '| parcial', {by.loc[i, 'TITULO']: (dem[i], int(by.loc[i, 'DISPONIBLE'])) for i in parcial})
