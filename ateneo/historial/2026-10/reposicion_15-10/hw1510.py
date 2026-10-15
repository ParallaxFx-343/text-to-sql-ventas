"""Repo 15/10 (5 sucursales) + Sugerido Halloween (completar a 10; chicos a 3; sin Juramento),
un PAEC por sucursal. Halloween = objetivo - (stock SII 15/10 + repo 14/10 en camino + repo 15/10), min 0.
Uso: python -P hw1510.py REPO_15.xlsx REPO_14.xlsx SII_15.xls config.json SALIDA"""
import sys, json, re, zipfile, unicodedata, pandas as pd
from pathlib import Path
from openpyxl import load_workbook
f15, f14, fsii, fcfg, out = sys.argv[1:6]; out=Path(out); out.mkdir(parents=True, exist_ok=True)
cfg=json.load(open(fcfg)); orden=cfg['orden_paec']+['DEVOTO']
CHICOS={'TRELEW','EZEIZA','PATIO BULLRICH','BAHIA BLANCA SHOPPING','PALERMO PORTAL'}; EXCL={'ATENEO JURAMENTO'}
MAPA={'NUEVO CENTRO':'NUEVO CENTRO SHOPPING','MAR DEL PLATA ALDREY':'MAR DEL PLATA PASEO ALDREY','SANTA FE':'SANTA FE Ribera Shopping'}
norm=lambda s: unicodedata.normalize('NFKD',s).encode('ascii','ignore').decode().upper().strip()
def bloques(f):
    cur=None; res={}
    for r in load_workbook(f).worksheets[0].iter_rows(values_only=True):
        r=r[:5]
        if r[0] is None and r[2] and r[1] is None: cur=MAPA.get(norm(r[2]),norm(r[2])); assert cur in orden,cur; res[cur]={}; continue
        if isinstance(r[0],int):
            i=int(r[3]); assert i not in res[cur]; res[cur][i]=int(r[1])
    return res
r15, r14 = bloques(f15), bloques(f14)
hw=[int(r[8]) for r in load_workbook(f15).worksheets[0].iter_rows(values_only=True) if isinstance(r[8],int)]
s=pd.read_excel(fsii); s=s[s['ID ARTICULO'].notna()].copy(); s['n']=pd.to_numeric(s.ISBN,errors='coerce'); by=s.set_index('n')
assert all(h in by.index for h in hw)
sor=lambda i: str(i).startswith('290')
lineas={}   # loc -> {isbn: [repo, hw]}
calc=[]
for loc in orden:
    if loc in EXCL: continue
    for i,v in r15.get(loc,{}).items():
        if not sor(i): lineas.setdefault(loc,{}).setdefault(i,[0,0])[0]=v
    obj=3 if loc in CHICOS else 10
    for h in hw:
        st=int(0 if pd.isna(by.loc[h,loc]) else by.loc[h,loc]); t14=r14.get(loc,{}).get(h,0); t15=r15.get(loc,{}).get(h,0)
        q=max(0,obj-(st+t14+t15)); calc.append((loc,h,obj,st,t14,t15,q))
        if q: lineas.setdefault(loc,{}).setdefault(h,[0,0])[1]=q
c=pd.DataFrame(calc,columns=['local','isbn','objetivo','stock_sii_15','repo_14_en_camino','repo_15','halloween'])
c.to_csv(out/'calculo_halloween.csv',sep=';',index=False,encoding='utf-8-sig')
# disponible
dem={}
for loc,d in lineas.items():
    for i,(a,b) in d.items(): dem[i]=dem.get(i,0)+a+b
disp=lambda i: int(by.loc[i,'DISPONIBLE'])
corto={i:(dem[i],disp(i)) for i in dem if 0<disp(i)<dem[i]}
hw_corto={i:x for i,x in corto.items() if i in hw}
print('halloween por titulo:', {str(by.loc[h,'TITULO'])[:25]:int(c[c.isbn==h].halloween.sum()) for h in hw})
print('no alcanza el disponible (total):', {str(by.loc[i,'TITULO'])[:30]:x for i,x in corto.items()})
T=['REPOSICION 15/10 + SUGERIDO HALLOWEEN - UN PAEC POR SUCURSAL',
   'Repo (Solar de la Abadía, Flores, MdP Paseo Aldrey, Mendoza Palmares, Santa Fe): cantidad = "Ventas", sin Libros Sorpresa.',
   'Halloween (13 títulos): completar a 10; Trelew, Ezeiza, Patio Bullrich, Bahía Blanca y Palermo Portal a 3; sin Ateneo Juramento.',
   '   Se descuenta lo que ya va en camino: repo del 14/10 (todavía no figura en el SII) y la repo de hoy.',
   'Marcas: [H] = Halloween · [R+H] = repo + Halloween en la misma línea · [SIN STOCK] = disponible 0 · [DISP x PARA y] = no alcanza para todos.',
   'Stock y disponible: SII 15/10 07:01. Orden de las sucursales: el de los TXT de PAEC (Devoto al final).','']
files=[]; tl=tu=th=tr=0
for loc in orden:
    if loc not in lineas: continue
    d=lineas[loc]; slug=re.sub(r'[^A-Z0-9]+','_',loc).strip('_')
    items=sorted(d.items(), key=lambda kv:(-(kv[1][0]+kv[1][1]), str(by.loc[kv[0],'TITULO'])))
    p=out/f'reposicion_{slug}.csv'; files.append(p)
    p.write_bytes(('\r\n'.join(f'{i};{a+b}' for i,(a,b) in items)+'\r\n').encode())
    ur=sum(a for a,b in d.values()); uh=sum(b for a,b in d.values())
    T+=['='*78, f'{loc:<50}{ur+uh:>6} u  {len(d)} títulos', f'   {p.name} | repo {ur} u | halloween {uh} u','='*78]
    for i,(a,b) in items:
        m='[R+H]' if a and b else ('[H]' if b else '')
        x='  [SIN STOCK]' if disp(i)<=0 else (f'  [DISP {disp(i)} PARA {corto[i][0]}]' if i in corto else '')
        T.append(f'   {a+b:>4}  {i}  {str(by.loc[i,"TITULO"])[:42]:<42} {m:<5}{x}' + (f'  (repo {a} + halloween {b})' if a and b else ''))
    T.append('')
    tl+=len(d); tu+=ur+uh; th+=uh; tr+=ur
T+=['='*78, f'TOTAL {len(lineas)} sucursales | {tl} líneas | {tu} u (repo {tr} + halloween {th})']
(out/'DETALLE_REPO_Y_HALLOWEEN_15-10.txt').write_text('\n'.join(T)+'\n',encoding='utf-8-sig')
with zipfile.ZipFile(out/'reposicion_y_halloween_15-10_por_sucursal.zip','w') as z:
    for p in files: z.write(p,p.name)
print(f'{len(lineas)} sucursales | {tl} lineas | {tu} u (repo {tr} + halloween {th})')
