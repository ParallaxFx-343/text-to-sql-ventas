"""Agrega los movimientos de sep-oct 2026 a la hoja 'Reposición x local' de Cadena Stock 26.27.
Cada movimiento va como fila nueva debajo de la ultima fila de su sucursal, con el formato de las
filas de detalle existentes. DEVOTO (sin grupo) se agrega al final.
Uso: python -P cargar_cadena.py ENTRADA.xlsx SALIDA.xlsx"""
import sys, copy, json
import os; sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from movimientos import mov
from openpyxl import load_workbook
fin, fout = sys.argv[1:3]
wb = load_workbook(fin); ws = wb['Reposición x local']
pad = json.load(open('/home/user/text-to-sql-ventas/ateneo/config.json'))['orden_paec'] + ['DEVOTO']
ALIAS = {'QUILMES': 'QUILMES CENTRO', 'REMEROS': 'REMEROS (NORDELTA)', 'BENDÚ': 'MAR DEL PLATA M. BENDU',
         'LOS GALLEGOS': 'MAR DEL PLATA LOS GALLEGOS', 'MAR DEL PLATA SHOP': 'MAR DEL PLATA PASEO ALDREY'}
# bloques de la hoja
bloques = {}; cur = None
for r in range(1, ws.max_row + 1):
    a = ws.cell(r, 1).value
    if a and not str(a).startswith('GRUPO') and r > 1:
        cur = ALIAS.get(a, a); assert cur in pad or cur == 'CANALES', a; bloques[cur] = [r, r]
    elif a and str(a).startswith('GRUPO'):
        cur = None
    elif cur and any(ws.cell(r, c).value is not None for c in range(2, 6)):
        bloques[cur][1] = r
faltan = [l for l in {m[0] for m in mov} if l not in bloques]
assert faltan == ['DEVOTO'], faltan
movs = {}
for l, c, d, p, u in sorted(mov, key=lambda m: (m[2], m[1])):
    movs.setdefault(l, []).append((c, d, p, u))
grupos = [str(m) for m in ws.merged_cells.ranges]
modelo = ws.max_row - 1 if False else None
det = 4   # fila modelo de detalle (sin sucursal en A)
def estilo(dst, src):
    for c in range(1, 6):
        s, t = ws.cell(src, c), ws.cell(dst, c)
        t._style = copy.copy(s._style)
# canales: la fila placeholder se reemplaza por el primer movimiento
r0, r1 = bloques['CANALES']
assert ws.cell(r0, 2).value == '— sin PAEC cargado —'
# insertar de abajo hacia arriba
orden = sorted([l for l in movs if l != 'DEVOTO'], key=lambda l: -bloques[l][1])
for l in orden:
    first, last = bloques[l]
    items = movs[l]
    if l == 'CANALES':
        c, d, p, u = items[0]
        for col, v in zip(range(2, 6), (c, d, p, u)): ws.cell(first, col).value = v
        ws.cell(first, 3).number_format = ws.cell(det, 3).number_format
        items = items[1:]
    n = len(items)
    if not n: continue
    ws.insert_rows(last + 1, n)
    for k, (c, d, p, u) in enumerate(items):
        r = last + 1 + k
        estilo(r, det if det < last + 1 else det + n)
        ws.cell(r, 1).value = None
        for col, v in zip(range(2, 6), (c, d, p, u)): ws.cell(r, col).value = v
    # mover rangos combinados de los encabezados de grupo que quedaron abajo
    nuevos = []
    for g in list(ws.merged_cells.ranges):
        if g.min_row > last:
            ws.merged_cells.ranges.remove(g) if hasattr(ws.merged_cells.ranges, 'remove') else ws.merged_cells.ranges.discard(g)
            nuevos.append((g.min_row + n, g.max_row + n, g.min_col, g.max_col))
    for a, b, c1, c2 in nuevos:
        ws.merge_cells(start_row=a, end_row=b, start_column=c1, end_column=c2)
# DEVOTO al final
r = ws.max_row + 1
cab = 3  # fila modelo de primera fila de sucursal
for k, (c, d, p, u) in enumerate(movs['DEVOTO']):
    estilo(r + k, cab if k == 0 else det)
    ws.cell(r + k, 1).value = 'DEVOTO' if k == 0 else None
    for col, v in zip(range(2, 6), (c, d, p, u)): ws.cell(r + k, col).value = v
wb.save(fout)
print('ok', sum(len(v) for v in movs.values()), 'movimientos;', len(movs), 'sucursales')
