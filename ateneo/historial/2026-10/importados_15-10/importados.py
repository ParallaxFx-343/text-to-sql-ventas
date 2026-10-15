"""Sabana de titulos importados (coediciones): entregados y liquidados por cliente.
Uso: python -P importados.py SABANA.xlsx SALIDA.xlsx"""
import sys, datetime as dt, pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
fin, fout = sys.argv[1:3]
s=pd.read_excel(fin,sheet_name='Sabana consolidada'); l=pd.read_excel(fin,sheet_name='Listado Importados',header=1)
assert set(s.CODIGO)==set(l.IDARTICULO)
assert ((s.ENTREGADA-s.VENTA-s.DEVOL+s.AJUSTE+s.TRANSF)==s.SALDO).all()
s['fliq']=pd.to_datetime(s.FECHAULTLIQ, unit='D', origin='1899-12-30')
c=s.groupby('CUENTA').agg(cliente=('CLIENTE','first'),vendedor=('NOMBRE_VENDEDOR','first'),
    sucursales=('SUCURSAL','nunique'),titulos=('CODIGO','nunique'),ent=('ENTREGADA','sum'),ven=('VENTA','sum'),
    dev=('DEVOL','sum'),aj=('AJUSTE','sum'),saldo=('SALDO','sum'),fliq=('fliq','max'),
    tit_sin=('VENTA',lambda v:(v==0).sum())).reset_index()
c['pct']=c.ven/c.ent; c['part']=c.ven/c.ven.sum()
c=c.sort_values(['ent','ven'],ascending=False)
sin=c[c.ven==0].sort_values('ent',ascending=False)
H=Font(bold=True,color='FFFFFF'); HF=PatternFill('solid',fgColor='1F4E78'); b=Side(style='thin',color='D9DCE1'); caja=Border(left=b,right=b,top=b,bottom=b)
def hoja(ws,titulo,cols,filas,anchos,fmts,total_cols):
    ws['A1']=titulo; ws['A1'].font=Font(bold=True,size=12)
    for j,t in enumerate(cols,1):
        x=ws.cell(3,j,t); x.font=H; x.fill=HF; x.border=caja; x.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for i,f in enumerate(filas,4):
        for j,v in enumerate(f,1):
            x=ws.cell(i,j,v); x.border=caja
            if fmts.get(j): x.number_format=fmts[j]
    u=3+len(filas); t=u+1
    ws.cell(t,2,'TOTAL').font=Font(bold=True)
    for j in total_cols:
        x=ws.cell(t,j,f'=SUM({L(j)}4:{L(j)}{u})'); x.font=Font(bold=True); x.border=caja
    for j,w in enumerate(anchos,1): ws.column_dimensions[L(j)].width=w
    ws.row_dimensions[3].height=32; ws.freeze_panes='C4'; ws.auto_filter.ref=f'A3:{L(len(cols))}{u}'
    return u,t
wb=Workbook()
# Resumen
r=wb.active; r.title='Resumen'
tot_e,tot_v=int(s.ENTREGADA.sum()),int(s.VENTA.sum())
datos=[('Clientes (cuentas)',len(c),'0'),('Títulos importados',s.CODIGO.nunique(),'0'),
       ('Ejemplares entregados',tot_e,'#,##0'),('Ejemplares liquidados (venta)',tot_v,'#,##0'),
       ('% liquidado sobre entregado',tot_v/tot_e,'0.0%'),('Devueltos',int(s.DEVOL.sum()),'#,##0'),
       ('Saldo en consignación',int(s.SALDO.sum()),'#,##0'),
       ('Clientes sin ningún ejemplar liquidado',len(sin),'0'),('   ejemplares entregados a esos clientes',int(sin.ent.sum()),'#,##0')]
r['A1']='Títulos importados por coediciones · entregados y liquidados por cliente'; r['A1'].font=Font(bold=True,size=13)
for i,(k,v,f) in enumerate(datos,3):
    r.cell(i,1,k); x=r.cell(i,2,v); x.number_format=f; x.font=Font(bold=True)
r.cell(14,1,'"Liquidados" = columna VENTA de la sábana. "% liquidado" = liquidados / entregados del cliente. '
       '"Participación" = qué parte del total liquidado de estos títulos explica el cliente.').font=Font(italic=True,size=9)
r.cell(15,1,'Un cliente = una cuenta; si tiene más de una sucursal, se suman. Fuente: Sabana consolidada (57 títulos del Listado Importados).').font=Font(italic=True,size=9)
r.column_dimensions['A'].width=44; r.column_dimensions['B'].width=14
# Por cliente
ws=wb.create_sheet('Por cliente')
cols=['Cuenta','Cliente','Vendedor','Suc.','Títulos','Entregados','Liquidados','% liquidado','Participación en lo liquidado','Devueltos','Saldo','Títulos sin venta','Última liquidación']
filas=[[x.CUENTA,x.cliente,x.vendedor,x.sucursales,x.titulos,int(x.ent),int(x.ven),x.pct,x.part,int(x.dev),int(x.saldo),int(x.tit_sin),x.fliq.to_pydatetime() if pd.notna(x.fliq) else None] for x in c.itertuples()]
u,t=hoja(ws,'Entregados y liquidados por cliente · ordenado por entregados',cols,filas,[8,44,24,6,8,11,11,10,13,10,9,10,12],
         {6:'#,##0',7:'#,##0',8:'0.0%',9:'0.00%',10:'#,##0',11:'#,##0',13:'DD/MM/YYYY'},[6,7,10,11])
ws.cell(t,8,f'=G{t}/F{t}').number_format='0.0%'; ws.cell(t,8).font=Font(bold=True)
# Sin liquidar
ws2=wb.create_sheet('Sin liquidar')
cols2=['Cuenta','Cliente','Vendedor','Suc.','Títulos entregados','Entregados','Devueltos','Saldo','Última liquidación']
filas2=[[x.CUENTA,x.cliente,x.vendedor,x.sucursales,x.titulos,int(x.ent),int(x.dev),int(x.saldo),x.fliq.to_pydatetime() if pd.notna(x.fliq) else 'nunca'] for x in sin.itertuples()]
hoja(ws2,'Clientes que no liquidaron ningún ejemplar de estos títulos · ordenado por entregados',cols2,filas2,[8,44,24,6,10,11,10,9,12],
     {6:'#,##0',7:'#,##0',8:'#,##0',9:'DD/MM/YYYY'},[6,7,8])
# Por titulo
t=s.groupby('CODIGO').agg(isbn=('ISBN','first'),titulo=('TITULO','first'),clientes=('CUENTA','nunique'),ent=('ENTREGADA','sum'),
    ven=('VENTA','sum'),dev=('DEVOL','sum'),saldo=('SALDO','sum'),
    cli_sin=('CUENTA',lambda x: x[s.loc[x.index].groupby(x).VENTA.transform('sum')==0].nunique())).reset_index()
ped=l.groupby('IDARTICULO').NROPEDIDOYENNY.apply(lambda x:' / '.join(str(v) for v in sorted(set(x))))
t['ped']=t.CODIGO.map(ped); t['pct']=t.ven/t.ent; t['part']=t.ven/t.ven.sum()
t=t.sort_values(['ent','ven'],ascending=False)
ws3=wb.create_sheet('Por título')
cols3=['ID','ISBN','Título','Clientes','Entregados','Liquidados','% liquidado','Participación en lo liquidado','Devueltos','Saldo','Clientes sin venta del título','Pedido Yenny']
filas3=[[int(x.CODIGO),str(x.isbn),x.titulo,int(x.clientes),int(x.ent),int(x.ven),x.pct,x.part,int(x.dev),int(x.saldo),int(x.cli_sin),x.ped] for x in t.itertuples()]
u3,t3=hoja(ws3,'Entregados y liquidados por título · ordenado por entregados',cols3,filas3,[8,15,44,9,11,11,10,13,10,9,12,26],
           {5:'#,##0',6:'#,##0',7:'0.0%',8:'0.00%',9:'#,##0',10:'#,##0'},[5,6,9,10])
ws3.cell(t3,7,f'=F{t3}/E{t3}').number_format='0.0%'; ws3.cell(t3,7).font=Font(bold=True)
print(t.head(5)[['titulo','clientes','ent','ven','pct','cli_sin']].to_string(index=False)); print('...'); print(t.sort_values('pct').head(5)[['titulo','ent','ven','pct']].to_string(index=False))
wb.save(fout)
print(len(c),'clientes | ent',tot_e,'ven',tot_v,f'{tot_v/tot_e:.1%}','| sin liquidar',len(sin),int(sin.ent.sum()))
print(sin[['CUENTA','cliente','vendedor','titulos','ent','dev','saldo','fliq']].to_string(index=False))
print(c.head(10)[['CUENTA','cliente','ent','ven','pct','part']].to_string(index=False))
