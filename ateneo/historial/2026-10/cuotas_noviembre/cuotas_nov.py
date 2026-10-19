"""Cuotas de novedades: carga las novedades del mes nuevo en la planilla del mes anterior.
Escribe titulo (fila 2), ISBN (3), tirada (20) y cronograma (21) en C en adelante; vacia los
valores del mes anterior (IDs fila 5, cuotas 6-16, columnas de novedades sobrantes y
reimpresiones). No toca formulas (17, 18, 23), rotulos ni formatos, salvo el formato de fecha
y el relleno amarillo "ya en depo" de la fila 21 en las columnas que se escriben.
Uso: python -P cuotas_nov.py ENTRADA.xlsx SALIDA.xlsx"""
import sys, datetime as dt, copy
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
NOV = [  # titulo, isbn, tirada, cronograma (fecha de ingreso)
    ('BRIGID', 9789500218535, 3500, '16/10 o 19/10'),
    ('AVENTURERAS', 9789500218610, 3000, dt.datetime(2026, 10, 26)),
    ('LA ESCLAVA DE QURTUBA', 9789500218641, 4000, dt.datetime(2026, 10, 26)),
    ('CADA NIÑO, UN MUNDO', 9789500218627, 3000, dt.datetime(2026, 10, 20)),
    ('ALICIA EN EL PAIS DE LAS MARAVILLAS (ED. COMENTADA)', 9789500218634, 4000, dt.datetime(2026, 10, 23)),
    ('UN CUENTO DE NAVIDAD - A CHRISTMAS CAROL (ED. TAPA DURA, BILINGÜE)', 9789500218504, 4000, dt.datetime(2026, 10, 23)),
    ('EL KYBALION', 9789500218566, 4500, dt.datetime(2026, 10, 15)),
    ('SOBRE RUEDAS. AL RESCATE', 9789500217699, 4000, dt.datetime(2026, 10, 23)),
    ('SOBRE RUEDAS. EN LA GRANJA', 9789500217705, 4000, dt.datetime(2026, 10, 23)),
]
def dv(i):
    t = str(i); return (10 - sum(int(d) * (1 if k % 2 == 0 else 3) for k, d in enumerate(t[:12])) % 10) % 10 == int(t[12])
assert all(dv(i) for _, i, _, _ in NOV)
fin, fout = sys.argv[1:3]
wb = load_workbook(fin); ws = wb['CUOTAS']
assert ws['C1'].value == 'NOVEDADES' and ws['O1'].value == 'REIMPRESIONES' and ws['B20'].value == 'TIRADA' and ws['B21'].value == 'Cronograma'
ult = max(c.column for c in ws[2] if c.value is not None)          # ultima columna con titulo (AD)
# vaciar valores del mes anterior (no formulas)
for col in range(CI('C'), ult + 1):
    for r in [2, 3, 5] + list(range(6, 17)) + [20, 21]:
        c = ws.cell(r, col)
        if not (isinstance(c.value, str) and c.value.startswith('=')):
            c.value = None
sin_relleno = copy.copy(ws['C21'].fill)
for k, (t, i, q, f) in enumerate(NOV):
    col = CI('C') + k
    ws.cell(2, col).value = t; ws.cell(3, col).value = i; ws.cell(20, col).value = q
    c = ws.cell(21, col); c.value = f; c.number_format = 'dd/mm'; c.fill = copy.copy(sin_relleno)
wb.save(fout)
print('ok', len(NOV), 'novedades en C..' + L(CI('C') + len(NOV) - 1), '| vaciado hasta', L(ult))
