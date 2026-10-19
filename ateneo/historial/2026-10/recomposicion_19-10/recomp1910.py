"""Recomposicion de fondo 19/10/2026: 1 ejemplar de cada titulo de la lista a los locales
que tienen 0 (SII), excluyendo Ezeiza, Patio Bullrich, Trelew, Devoto y Juramento.
Fuera: los 29 titulos de la hoja "Libro sorpresa" (pedido aparte) y las filas ocultas de la lista.
Uso: python -P recomp1910.py LISTA.xlsx EXCLUIR.xlsx SII.xls config.json amba.txt SALIDA"""
import sys, json, re, zipfile, pandas as pd
from pathlib import Path
from openpyxl import load_workbook
flista, fex, fsii, fcfg, famba, out = sys.argv[1:7]; out=Path(out); out.mkdir(parents=True, exist_ok=True)
EXCL={'EZEIZA','PATIO BULLRICH','TRELEW','DEVOTO','ATENEO JURAMENTO'}
pad=json.load(open(fcfg))['orden_paec']+['DEVOTO']; locs=[p for p in pad if p not in EXCL]
AMBA=open(famba,encoding='utf-8').read().split('\n')
ws=load_workbook(flista).active
lista=[]; ocultas=[]
for r in range(4,ws.max_row+1):
    v=ws.cell(r,1).value
    if v is None: continue
    x=(int(str(v).strip()),ws.cell(r,2).value,ws.cell(r,3).value,r)
    (ocultas if ws.row_dimensions[r].hidden else lista).append(x)
sor={int(r[0]):r[6] for r in load_workbook(fex)['Libro sorpresa'].iter_rows(min_row=5,values_only=True) if r[0]}
s=pd.read_excel(fsii); s=s[s['ID ARTICULO'].notna()].copy(); s['n']=pd.to_numeric(s.ISBN,errors='coerce'); by=s.set_index('n')
assert all(p in s.columns for p in locs)
vistos=set(); tit=[]; dup=[]; fuera_sor=[]
for i,idart,t,r in lista:
    if i in vistos: dup.append((i,t,r)); continue
    vistos.add(i)
    if i in sor: fuera_sor.append((i,t)); continue
    assert i in by.index, (i,t); assert int(by.loc[i,'ID ARTICULO'])==idart, (i,idart)
    tit.append(i)
filas=[]
for i in tit:
    for l in locs:
        st=by.loc[i,l]; st=0 if pd.isna(st) else int(st)
        if st<=0: filas.append((i,l))
rep=pd.DataFrame(filas,columns=['isbn','local']); rep['cant']=1
dem=rep.groupby('isbn').cant.sum()
disp={i:int(by.loc[i,'DISPONIBLE']) for i in dem.index}
corto={i:(int(dem[i]),disp[i]) for i in dem.index if disp[i]<dem[i]}
T=lambda i: str(by.loc[i,'TITULO'])
ID=lambda i: int(by.loc[i,'ID ARTICULO'])
orden_t=sorted(dem.index,key=lambda i:T(i))
# 1) PAEC unico
(out/'PAEC_RECOMPOSICION_19-10.csv').write_bytes(('\r\n'.join(f'{i};{int(dem[i])}' for i in orden_t)+'\r\n').encode())
cab=['RECOMPOSICION DE FONDO - 19/10/2026 - 1 ejemplar a los locales con stock 0',
     f'Stock y disponible: SII 19/10 07:01. {len(dem)} títulos | {rep.local.nunique()} locales | {len(rep)} unidades.',
     'Excluidos: Ezeiza, Patio Bullrich, Trelew, Devoto y Ateneo Juramento.',
     f'Fuera de la lista: los {len(fuera_sor)} títulos de "Libro sorpresa" y las {len(ocultas)} filas ocultas.','']
if corto: cab.insert(4,'OJO, disponible que no alcanza: '+'; '.join(f'{T(i)} pide {a}, hay {b}' for i,(a,b) in corto.items()))
L=list(cab)
for k,i in enumerate(orden_t,1):
    ls=sorted(rep[rep.isbn==i].local, key=lambda l:(l not in AMBA,l))
    L+=['='*70,f'{k}. {T(i)[:52]:<52}{int(dem[i]):>5} u',f'   ISBN {i}   ID {ID(i)}   en {len(ls)} locales   disp {disp[i]}','='*70]
    L+=[f'      1  {l}' for l in ls]; L.append('')
L+=['='*70,f'TOTAL {len(rep)} u']
(out/'DESGLOSE_x_TITULO_RECOMPOSICION_19-10.txt').write_text('\n'.join(L)+'\n',encoding='utf-8-sig')
# 2) por sucursal: TXT (AMBA primero) + CSV por local
M=list(cab); M[0]=M[0]+' - POR SUCURSAL (primero AMBA, después interior)'
files=[]
for zona in ('AMBA','INTERIOR'):
    zl=sorted(l for l in rep.local.unique() if (l in AMBA)==(zona=='AMBA'))
    M+=['#'*70,f'# {zona}: {len(zl)} locales, {int(rep[rep.local.isin(zl)].cant.sum())} u','#'*70,'']
    for l in zl:
        ii=sorted(rep[rep.local==l].isbn, key=T)
        slug=re.sub(r'[^A-Z0-9]+','_',l).strip('_'); p=out/f'recomposicion_{slug}.csv'; files.append(p)
        p.write_bytes(('\r\n'.join(f'{i};1' for i in ii)+'\r\n').encode())
        M+=['='*70,f'{l:<56}{len(ii):>5} u',f'   {p.name}','='*70]+[f'      1  {i}  {T(i)[:48]}' for i in ii]+['']
M+=['='*70,f'TOTAL {len(rep)} u']
(out/'DESGLOSE_x_SUCURSAL_RECOMPOSICION_19-10.txt').write_text('\n'.join(M)+'\n',encoding='utf-8-sig')
with zipfile.ZipFile(out/'recomposicion_19-10_por_sucursal.zip','w') as z:
    for p in files: z.write(p,p.name)
for p in files: p.unlink()
# 3) ocultas (no van) con lo que serian
O=['FILAS OCULTAS DE LA LISTA (NO SE MANDARON) - por si hay que sumarlas','SII 19/10: locales en 0 (de los 57) y disponible','']
for i,idart,t,r in ocultas:
    c=int((by.loc[i,locs].fillna(0)<=0).sum()) if i in by.index else None
    O.append(f'fila {r:>3}  {i}  {str(t)[:44]:<44}  locales en 0: {c:>3}  disp {int(by.loc[i,"DISPONIBLE"]) if i in by.index else "?"}')
(out/'NO_VAN_filas_ocultas.txt').write_text('\n'.join(O)+'\n',encoding='utf-8-sig')
rep.to_csv(out/'detalle_recomposicion.csv',sep=';',index=False,encoding='utf-8-sig')
print(f'{len(lista)} visibles ({len(dup)} duplicadas: {[(t,r) for _,t,r in dup]}), {len(ocultas)} ocultas, {len(fuera_sor)} sorpresa fuera')
print(f'{len(tit)} titulos evaluados | {len(dem)} con algun local en 0 | {len(rep)} u | {rep.local.nunique()} locales | corto {corto}')
print('por local min/max', rep.groupby('local').cant.sum().min(), rep.groupby('local').cant.sum().max())
print(rep.groupby('local').cant.sum().sort_values(ascending=False).head(6).to_dict())
