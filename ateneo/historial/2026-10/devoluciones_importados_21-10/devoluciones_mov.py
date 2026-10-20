"""Candidatos a devolución de los 5 títulos (Banderas, Gatito, Dino, Stickers Dinosaurios y Granja)
a partir de los MOVIMIENTOS de consignación (envíos, liquidaciones, devoluciones, transferencias).
Uso: python -P devoluciones_mov.py SALIDA.xlsx MOVIMIENTOS.xlsx SABANA.xlsx [SABANA2.xlsx ...]
La sábana solo aporta el vendedor: su ENTREGADA/VENTA cuenta únicamente los envíos que siguen
abiertos (con saldo), así que subestima lo liquidado de quien liquida remitos enteros (Cúspide).
Criterio sobre lo acumulado: Prioridad alta = stock >= 20 y liquidó < 25 % de lo enviado;
Pedir = stock >= 10 y < 25 %; Revisar = stock >= 10 sin ninguna liquidación o con 25-50 %; resto No."""
import os, sys, pandas as pd
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from parse_mov import parse
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
fout, fmov, fsab = sys.argv[1], sys.argv[2], sys.argv[3:]
T = {744138: 'Banderas del mundo', 747643: 'Gatito', 747645: 'Dino', 747646: 'Stickers Dinosaurios', 747647: 'Stickers Granja'}
h, m = parse(fmov)
assert set(h.Codigo) == set(T)
m['t'] = m.Mov.str.strip().map({'ENVIO CONSIGNACION': 'Enviado', 'LIQUIDACION DE CONSIGNACION': 'Liquidado',
                                 'DEVOLUCIONES CLIENTES EXTERNOS': 'Devuelto', 'CNSG ATENEO: TRANSFERENCIA DE': 'Transf.'})
assert m.t.notna().all()
assert (m.groupby('Codigo').Ctd.sum() == h.set_index('Codigo').Restante).all()   # cuadre contra el encabezado
sab = pd.concat([pd.read_excel(f, sheet_name=0).dropna(subset=['CODIGO']) for f in fsab])
vend = sab.groupby('CUENTA').NOMBRE_VENDEDOR.first()
k = ['Cliente', 'Nombre', 'Sucursal']
def agg(df, keys):
    p = df.pivot_table(index=keys, columns='t', values='Ctd', aggfunc='sum', fill_value=0)
    for c in ['Enviado', 'Liquidado', 'Devuelto', 'Transf.']:
        if c not in p: p[c] = 0
    p['Stock'] = p[['Enviado', 'Liquidado', 'Devuelto', 'Transf.']].sum(axis=1)
    p[['Liquidado', 'Devuelto']] = -p[['Liquidado', 'Devuelto']]
    f = df.groupby(keys).agg(**{'Primer envío': ('Fecha', lambda s: s[df.loc[s.index, 't'] == 'Enviado'].min()),
                                'Últ. liq.': ('Fecha', lambda s: s[df.loc[s.index, 't'] == 'Liquidado'].max())})
    p = p.join(f).reset_index()
    p['% liquidado'] = p.Liquidado / p.Enviado
    return p
g = agg(m, k)
st = m.assign(Título=m.Codigo.map(T)).pivot_table(index=k, columns='Título', values='Ctd', aggfunc='sum', fill_value=0)
st = st[[t for t in T.values()]].add_prefix('Stock ').reset_index()
g = g.merge(st, on=k)
g['Vendedor'] = g.Cliente.map(vend)
g['Sucursales de la cuenta'] = g.Cliente.map(g[g.Stock > 0].groupby('Cliente').Sucursal.nunique()).fillna(0).astype(int)
def sug(r):
    if r.Stock >= 10 and pd.isna(r['Últ. liq.']): return 'Revisar (nunca liquidó)'
    if r.Stock >= 20 and r['% liquidado'] < .25: return 'Prioridad alta'
    if r.Stock >= 10 and r['% liquidado'] < .25: return 'Pedir'
    if r.Stock >= 10 and r['% liquidado'] < .5: return 'Revisar'
    return 'No'
g['Sugerencia'] = g.apply(sug, axis=1)
orden = {'Prioridad alta': 0, 'Pedir': 1, 'Revisar (nunca liquidó)': 2, 'Revisar': 3, 'No': 4}
g = g[g.Stock > 0].sort_values(['Sugerencia', 'Stock'], key=lambda c: c.map(orden) if c.name == 'Sugerencia' else -c)
g = g.rename(columns={'Cliente': 'Cuenta', 'Nombre': 'Cliente'})
cols = ['Sugerencia', 'Cuenta', 'Cliente', 'Sucursal', 'Sucursales de la cuenta', 'Vendedor', 'Primer envío', 'Últ. liq.',
        'Enviado', 'Liquidado', 'Devuelto', 'Transf.', 'Stock', '% liquidado'] + [c for c in g.columns if c.startswith('Stock ')]
g = g[cols]
m['Título'] = m.Codigo.map(T)
det = agg(m, k + ['Codigo', 'Título']).rename(columns={'Cliente': 'Cuenta', 'Nombre': 'Cliente', 'Codigo': 'ID'})
det = det[det.Enviado > 0].sort_values('Stock', ascending=False)
tot = agg(m.assign(x=1), ['Codigo', 'Título']).rename(columns={'Codigo': 'ID'})
tot.insert(2, 'Clientes', tot.ID.map(m.groupby('Codigo').Cliente.nunique()))
cu = m.assign(Grupo=m.Cliente.map(lambda c: 'Cúspide' if c == 10842 else 'Resto de los clientes'))
cmp_ = agg(cu, ['Grupo', 'Título'])
res = g.groupby('Sugerencia', sort=False).agg(Puntos=('Cuenta', 'size'), Stock=('Stock', 'sum')).reset_index()
movs = m[['Cliente', 'Nombre', 'Sucursal', 'Título', 'Codigo', 'Mov', 'Comprobante', 'Factura', 'Ctd', 'Restante', 'Fecha', 'Obs']]
with pd.ExcelWriter(fout, engine='openpyxl') as w:
    g.to_excel(w, sheet_name='Por cliente', index=False)
    det.to_excel(w, sheet_name='Por cliente y título', index=False)
    tot.to_excel(w, sheet_name='Totales', index=False)
    res.to_excel(w, sheet_name='Totales', index=False, startrow=len(tot) + 3)
    cmp_.to_excel(w, sheet_name='Cúspide vs resto', index=False)
    movs.to_excel(w, sheet_name='Movimientos', index=False)
wb = load_workbook(fout)
col = {'Prioridad alta': 'F8CBAD', 'Pedir': 'FCE4D6', 'Revisar (nunca liquidó)': 'FFF2CC', 'Revisar': 'FFF2CC'}
W = {'Cliente': 40, 'Nombre': 40, 'Sucursal': 32, 'Vendedor': 22, 'Sugerencia': 22, 'Título': 22, 'Grupo': 22, 'Mov': 30, 'Comprobante': 20, 'Factura': 16, 'Obs': 24}
for ws in wb:
    for c in ws[1]: c.font = Font(bold=True); c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    for i, hd in enumerate([c.value for c in ws[1]], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = W.get(hd, 12)
        for r in range(2, ws.max_row + 1):
            c = ws.cell(r, i)
            if hd == '% liquidado': c.number_format = '0%'
            elif hd in ('Primer envío', 'Últ. liq.', 'Fecha'): c.number_format = 'DD/MM/YYYY'
    if ws.title not in ('Totales',): ws.auto_filter.ref = ws.dimensions
ws = wb['Por cliente']
for r in range(2, ws.max_row + 1):
    f = col.get(ws.cell(r, 1).value)
    if f:
        for c in range(1, ws.max_column + 1): ws.cell(r, c).fill = PatternFill('solid', fgColor=f)
wb.save(fout)
pd.set_option('display.width', 250)
print(tot.to_string()); print(cmp_.to_string()); print(res.to_string())
print(g[g.Sugerencia == 'Prioridad alta'][['Cliente', 'Vendedor', 'Primer envío', 'Enviado', 'Liquidado', 'Stock', '% liquidado']].to_string())
print(g[g.Cuenta == 10842].T.to_string())
