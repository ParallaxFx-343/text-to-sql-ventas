import sys, json, pandas as pd
from openpyxl import load_workbook
f, cfg = sys.argv[1:3]
pad=json.load(open(cfg))['orden_paec']+['DEVOTO']
ws=load_workbook(f).worksheets[0]
rows=[r[:5] for r in ws.iter_rows(values_only=True)]   # solo A-E: a la derecha puede venir otra tabla
blocks=[]; cur=None
for i,r in enumerate(rows,1):
    if r[0] is None and r[2] and all(x is None for k,x in enumerate(r) if k!=2):
        cur={'local':r[2],'fila':i,'lineas':[]}; blocks.append(cur); continue
    if r[0]=='Orden': continue
    if isinstance(r[0],(int,float)) and cur is not None:
        cur['lineas'].append((i,)+tuple(r[:5]))
    elif any(x is not None for x in r) and i>3:
        print('fila rara',i,r)
print(len(blocks),'bloques')
for b in blocks:
    L=b['lineas']; print(f"{b['local']:<30} fila {b['fila']:>3} | {len(L):>3} lineas | {sum(x[2] for x in L):>4} u | orden 1..{max(x[1] for x in L)} ok={[x[1] for x in L]==list(range(1,len(L)+1))}")
import pickle; pickle.dump(blocks,open('r2710/blocks.pkl','wb'))
