"""Sin SII: mapea cada bloque al nombre del padrón por nombre (sin ambigüedad en esta planilla) y hace los
controles que no necesitan SII: ISBN repetidos, dígito verificador, Ventas <= 0, Stock vacío, Sorpresa.
Uso: python -P repo2710_map.py r2710/blocks.pkl config.json excluir.xlsx"""
import sys, json, pickle, unicodedata
from openpyxl import load_workbook
blocks = pickle.load(open(sys.argv[1], 'rb'))
pad = json.load(open(sys.argv[2]))['orden_paec'] + ['DEVOTO']
EXSOR = {int(r[0]) for r in load_workbook(sys.argv[3])['Libro sorpresa'].iter_rows(min_row=5, values_only=True) if r[0]}
norm = lambda s: unicodedata.normalize('NFKD', s).encode('ascii', 'ignore').decode().upper().strip()
ALIAS = {'BAHIA BLANCA': 'BAHIA BLANCA SHOPPING', 'BENDU': 'MAR DEL PLATA M. BENDU'}
def dv(i):
    s = str(i); return len(s) == 13 and (10 - sum(int(d) * (1 if k % 2 == 0 else 3) for k, d in enumerate(s[:12])) % 10) % 10 == int(s[12])
for b in blocks:
    n = norm(b['local']); n = ALIAS.get(n, n)
    assert n in pad, n
    b['suc'] = n
    isb = [int(x[4]) for x in b['lineas']]
    print(f"{b['local']:<16} -> {n:<24} dup {len(isb)-len(set(isb))} dvmal {[i for i in isb if not dv(i)]} "
          f"ventas<=0 {[x[1] for x in b['lineas'] if not x[2] or x[2] <= 0]} stockvacio {[x[1] for x in b['lineas'] if x[5] is None]} "
          f"sorpresa {[x[1] for x in b['lineas'] if str(x[4]).startswith('290')]} lista {[(x[1], x[3]) for x in b['lineas'] if int(x[4]) in EXSOR]}")
pickle.dump(blocks, open(sys.argv[1], 'wb'))
