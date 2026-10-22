"""Servicio de Novedades noviembre 2026: escribe las cuotas de cada vendedor en su grilla y pone 0 en los
importados (Sobre ruedas x2) para los clientes de la solapa "Clientes a no enviar importados".
Solo se escriben valores: no se tocan colores, formatos, encabezados, filas ni columnas.
Uso: python -P cuotas_vendedores.py CARPETA_ZIP CARPETA_SALIDA"""
import sys, os, glob, openpyxl
from openpyxl.utils import column_index_from_string as CI
src, out = sys.argv[1:3]; os.makedirs(out, exist_ok=True)
cf = openpyxl.load_workbook(glob.glob(os.path.join(src, '*Cuotas novedades*'))[0], data_only=True)['CUOTAS']
NOVC = [c for c in range(3, 12) if cf.cell(3, c).value]                  # C..K = novedades
ISBN = [int(cf.cell(3, c).value) for c in NOVC]
CUO = {cf.cell(r, 2).value.strip(): [cf.cell(r, c).value for c in NOVC] for r in range(6, 16)}
IMPORT = {9789500217699, 9789500217705}            # SOBRE RUEDAS (coedición importada de China)
sab = openpyxl.load_workbook(glob.glob(os.path.join(src, '*Sabana*'))[0], read_only=True)['Clientes a no enviar importados']
NOENV = {int(r[0]): (r[1], r[2]) for r in sab.iter_rows(min_row=2, values_only=True) if r[0] and r[6] == 'No enviar Importados'}
# archivo, hoja, [(fila de cuotas, clave en la planilla de cuotas)]
GRILLAS = [('EGISTI', 'Hoja1', [(67, 'EGISTI, CLAUDIO')]),
           ('FUCITO SERVICIO', 'Hoja1', [(55, 'FUCITO  generales'), (56, 'FUCITO  Cúspide'), (57, 'FUCITO  Queleer'), (58, 'FUCITO  Santa Fe')]),
           ('FUCITO Queleer', 'Hoja1', [(20, 'FUCITO  Queleer')]),
           ('FUCITO Santa Fe', 'GRILLA', [(13, 'FUCITO  Santa Fe')]),
           ('HOUNDAKJIAN', 'Hoja1', [(29, 'HOUNDAKJIAN, PABLO')]),
           ('MARTIRENA', 'Hoja1', [(60, 'MARTIRENA, EZEQUIEL')]),
           ('NAVARRO', 'Hoja1', [(52, 'NAVARRO, GONZALO')]),
           ('RODRIGUEZ', 'Hoja1', [(34, 'RODRIGUEZ, GONZALO')]),
           ('SALICE', 'Hoja1', [(29, 'SÁLICE, MARCELO')])]
encontrados = {}
for pref, hoja, filas in GRILLAS:
    f = [p for p in glob.glob(os.path.join(src, f'{pref}*NOVIEMBRE 26.xlsx'))]; assert len(f) == 1, (pref, f)
    wb = openpyxl.load_workbook(f[0]); ws = wb[hoja]
    cols = {int(ws.cell(3, c).value): c for c in range(10, 19)}
    assert list(cols) == ISBN, (pref, list(cols))                       # mismas novedades y mismo orden
    for r, clave in filas:
        assert all(ws.cell(r, c).value is None for c in cols.values()), (pref, r)
        assert ws.cell(r, cols[ISBN[0]]).number_format == '0', (pref, r)  # la fila formateada de cuotas
        lab = ws.cell(r, 9).value
        assert lab is None or lab.strip() == clave.strip(), (pref, r, lab)
        for i, c in zip(ISBN, CUO[clave]): ws.cell(r, cols[i]).value = c
    last = max(r for r in range(6, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, (int, float)))
    ceros = []
    for r in range(6, last + 1):
        cta = ws.cell(r, 1).value
        if isinstance(cta, (int, float)) and int(cta) in NOENV:
            for i in IMPORT:
                assert ws.cell(r, cols[i]).value is None
                ws.cell(r, cols[i]).value = 0
            ceros.append((r, int(cta), ws.cell(r, 2).value)); encontrados.setdefault(int(cta), []).append(pref)
    wb.save(os.path.join(out, os.path.basename(f[0])))
    print(f'{pref:<16} cuotas en fila(s) {[r for r, _ in filas]} | importados en 0: {len(ceros)} filas', [f'{c} {n}' for _, c, n in ceros])
# Cúspide: novedades en filas, columna "Cuotas"
f = glob.glob(os.path.join(src, '*CUSPIDE*'))[0]; wb = openpyxl.load_workbook(f); ws = wb['El Ateneo - Ampliación locales']
cq = next(c for c in range(1, ws.max_column + 1) if str(ws.cell(4, c).value).strip().lower() == 'cuotas')
filas = {int(ws.cell(r, 1).value): r for r in range(5, 14)}
assert list(filas) == ISBN
for i, v in zip(ISBN, CUO['FUCITO  Cúspide']):
    assert ws.cell(filas[i], cq).value is None; ws.cell(filas[i], cq).value = v
wb.save(os.path.join(out, os.path.basename(f)))
print('CUSPIDE          cuotas en columna', ws.cell(4, cq).coordinate[:-1], 'filas 5-13 | 10842 no está en la lista de no enviar:', 10842 not in NOENV)
falt = {c: v for c, v in NOENV.items() if c not in encontrados}
print(f'\nNo enviar importados: {len(NOENV)} cuentas; en las grillas {len(encontrados)}; no aparecen en ninguna grilla {len(falt)}:')
for c, (n, v) in sorted(falt.items(), key=lambda x: x[1][1]): print('  ', c, n, '-', v)
