#!/usr/bin/env python3
"""Une los pedidos de un cliente (PEDIDO_* y REPO_*) en un solo archivo de carga.

Los clientes mandan un Excel por sucursal y por tipo de pedido, con el
nombre `{PEDIDO|REPO}_{CUENTA}_{CLIENTE}_SUC._{SUCURSAL}.xlsx` y una hoja
Hoja1 cuyo encabezado (Isbn, Cod., Titulo, Autor, Editorial, Cantidad, Tipo)
no esta siempre en la misma fila.

Salida, segun la convencion de CLAUDE.md ("Archivos de pedido por cliente"):
    {VENDEDOR}_-_{CLIENTE}_-_{CUENTA}.xlsx
      Hoja1   ISBN + Cantidad, sin encabezado, datos desde la fila 7
      Sheet1  ISBN + Descripcion + Cantidad, encabezado en fila 1

Controles:
  - digito verificador de cada ISBN (los clientes suelen mandar 978 + ISBN-10
    con el verificador viejo); se corrige y se avisa
  - todas las lineas de la misma cuenta y del mismo tipo (no mezclar
    consignacion con firme)
  - el mismo titulo con dos ISBN distintos (otra edicion)
  - cuadre: suma de los archivos de entrada = suma del archivo unido
  - con --sii, que cada ISBN exista y cuanto hay disponible

Uso:
    python pedido_cliente.py PEDIDO_*.xlsx REPO_*.xlsx --vendedor NAVARRO --salida ./out
    python pedido_cliente.py *.xlsx --vendedor X --reemplazar 9789500213653=9789500217514
"""
import argparse
import re
import sys
from collections import OrderedDict
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from isbn import resolver  # noqa: E402

PATRON = re.compile(r'(PEDIDO|REPO)_(\d+)_(.+?)_SUC\._(.+)\.xlsx$', re.I)


def leer(ruta):
    """Devuelve (meta, lineas) de un archivo de pedido del cliente."""
    nombre = re.sub(r'^[0-9a-f]{8}-', '', Path(ruta).name)
    m = PATRON.search(nombre)
    if not m:
        raise SystemExit(f'{nombre}: el nombre no sigue el patron PEDIDO|REPO_CUENTA_CLIENTE_SUC._X')
    meta = {'archivo': nombre, 'tipo_pedido': m.group(1).upper(), 'cuenta': m.group(2),
            'cliente': m.group(3).upper(), 'sucursal': m.group(4).upper()}

    ws = load_workbook(ruta, data_only=True)['Hoja1']
    fila_cab = None
    for fila in ws.iter_rows():
        if str(fila[0].value or '').strip().lower() == 'isbn':
            fila_cab = fila[0].row
            cab = [str(c.value or '').strip().lower() for c in fila]
            break
    if fila_cab is None:
        raise SystemExit(f'{nombre}: no encuentro el encabezado "Isbn" en Hoja1')

    def col(*claves):
        for i, c in enumerate(cab):
            if any(c.startswith(k) for k in claves):
                return i
        return None
    ci, cc, ct, cq, ctp = col('isbn'), col('cód', 'cod'), col('tít', 'tit'), col('cant'), col('tipo')

    lineas = []
    for fila in ws.iter_rows(min_row=fila_cab + 1, values_only=True):
        if fila[ci] is None or str(fila[ci]).strip() == '':
            continue
        original = str(fila[ci]).strip()
        if isinstance(fila[ci], float):
            original = str(int(fila[ci]))
        isbn, metodo = resolver(original)
        lineas.append({**meta, 'isbn_original': original, 'isbn': isbn, 'metodo': metodo,
                       'cod': fila[cc] if cc is not None else None,
                       'titulo': str(fila[ct]).strip() if ct is not None else '',
                       'cantidad': int(fila[cq]),
                       'tipo': str(fila[ctp]).strip() if ctp is not None else ''})
    return meta, lineas


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('archivos', nargs='+')
    ap.add_argument('--vendedor', default='VENDEDOR')
    ap.add_argument('--salida', default='.')
    ap.add_argument('--reemplazar', action='append', default=[],
                    help='ISBN_VIEJO=ISBN_NUEVO, decidido por Franco (se puede repetir)')
    ap.add_argument('--sii', help='SII para confirmar existencia y disponible')
    args = ap.parse_args()

    todas = []
    # Ordena por nombre real (sin el prefijo que agrega la subida): PEDIDO antes que REPO.
    lotes = [leer(ruta) for ruta in args.archivos]
    for meta, lineas in sorted(lotes, key=lambda x: x[0]['archivo']):
        todas += lineas

    reemplazos = dict(r.split('=') for r in args.reemplazar)
    for l in todas:
        if l['isbn'] in reemplazos:
            l['metodo'] = f"reemplazado ({l['isbn']} -> {reemplazos[l['isbn']]})"
            l['isbn'] = reemplazos[l['isbn']]

    # --- controles de consistencia ---
    cuentas = {l['cuenta'] for l in todas}
    clientes = {l['cliente'] for l in todas}
    tipos = {l['tipo'] for l in todas}
    if len(cuentas) > 1 or len(clientes) > 1:
        raise SystemExit(f'*** mezcla cuentas o clientes: {cuentas} {clientes}')
    if len(tipos) > 1:
        print(f'*** ATENCION: mezcla tipos de pedido {tipos}; revisar antes de unir')

    print(f'{"ARCHIVO":<44}{"LIN":>4}{"UNID":>6}')
    for nombre in OrderedDict.fromkeys(l['archivo'] for l in todas):
        ls = [l for l in todas if l['archivo'] == nombre]
        print(f'{nombre:<44}{len(ls):>4}{sum(l["cantidad"] for l in ls):>6}')

    corregidos = [l for l in todas if l['metodo'] not in ('ya era ISBN-13', '978 + codigo')]
    if corregidos:
        print('\nISBN corregidos:')
        for l in corregidos:
            print(f"  {l['isbn_original']:<18} -> {l['isbn']}  {l['metodo']}  "
                  f"({l['sucursal']}, {l['titulo']})")

    por_titulo = {}
    for l in todas:
        clave = re.sub(r'\s+', ' ', l['titulo'].upper())
        por_titulo.setdefault(clave, set()).add(l['isbn'])
    dobles = {t: s for t, s in por_titulo.items() if len(s) > 1}
    if dobles:
        print('\nmismo titulo con ISBN distinto (otra edicion?):')
        for t, s in dobles.items():
            print(f'  {t}: {", ".join(sorted(s))}')

    # --- union por ISBN, en orden de aparicion ---
    unido = OrderedDict()
    for l in todas:
        u = unido.setdefault(l['isbn'], {'isbn': l['isbn'], 'titulo': l['titulo'],
                                          'cantidad': 0, 'detalle': []})
        u['cantidad'] += l['cantidad']
        u['detalle'].append(f"{l['sucursal']} {l['tipo_pedido']} {l['cantidad']}")

    if args.sii:
        sii = pd.read_excel(args.sii)
        sii['ISBN'] = pd.to_numeric(sii['ISBN'], errors='coerce')
        sii = sii.dropna(subset=['ISBN']).set_index('ISBN')
        print('\ncontra el SII:')
        for u in unido.values():
            n = int(u['isbn'])
            if n not in sii.index:
                print(f"  *** {u['isbn']} NO esta en el SII ({u['titulo']})")
                continue
            f = sii.loc[n]
            disp = int(f.get('DISPONIBLE', 0) or 0)
            marca = 'OK ' if disp >= u['cantidad'] else '***'
            print(f"  {marca} {u['isbn']}  ID {int(f['ID ARTICULO']):<7} pide {u['cantidad']:>3}  "
                  f"disp {disp:>6}  {str(f['TITULO']).strip()[:40]}")

    entrada = sum(l['cantidad'] for l in todas)
    salida_total = sum(u['cantidad'] for u in unido.values())
    print(f'\ncuadre: entrada {entrada} = unido {salida_total} '
          f'{"OK" if entrada == salida_total else "*** NO CUADRA ***"}  '
          f'({len(todas)} lineas -> {len(unido)} ISBN)')

    # --- archivo de carga ---
    cliente, cuenta = clientes.pop(), cuentas.pop()
    wb = Workbook()
    h1 = wb.active
    h1.title = 'Hoja1'
    for i, u in enumerate(unido.values(), start=7):
        h1.cell(i, 1, int(u['isbn'])).number_format = '0'
        h1.cell(i, 2, u['cantidad'])
    s1 = wb.create_sheet('Sheet1')
    s1.append(['ISBN', 'Descripción', 'Cantidad'])
    for u in unido.values():
        s1.append([int(u['isbn']), u['titulo'], u['cantidad']])
        s1.cell(s1.max_row, 1).number_format = '0'
    for ws in (h1, s1):
        ws.column_dimensions['A'].width = 16
    s1.column_dimensions['B'].width = 46

    destino = Path(args.salida)
    destino.mkdir(parents=True, exist_ok=True)
    archivo = destino / f'{args.vendedor.upper()}_-_{cliente}_-_{cuenta}.xlsx'
    wb.save(archivo)

    print(f'\n{"ISBN":<15}{"CANT":>5}  {"TITULO":<42}DETALLE')
    for u in unido.values():
        print(f"{u['isbn']:<15}{u['cantidad']:>5}  {u['titulo'][:40]:<42}{' + '.join(u['detalle'])}")
    print(f'\nescrito {archivo}')


if __name__ == '__main__':
    main()
