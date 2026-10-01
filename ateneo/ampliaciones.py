#!/usr/bin/env python3
"""Arma los archivos de carga a SILOMA de las ampliaciones de un vendedor.

Entrada: los `AMPLIACION - {CLIENTE} {CUENTA} {fecha} EL ATENEO.xls|xlsx|ods`
que manda el vendedor, uno por cliente. Vienen con encabezados distintos
(CUENTA/CLIENTE/CODIGO/ISBN/TITULO/PVP/AMPLIACION, o CODIGO = ISBN sin cuenta,
o con SALDO...) y a veces con filas de relleno sin ISBN.

Salida, un archivo por cliente (convencion "Archivos de pedido por cliente",
solo Hoja1, como pidio Franco para las ampliaciones):
    {VENDEDOR}_-_{CLIENTE}_-_{CUENTA}.xlsx
      Hoja1: ISBN (entero, formato 0) + cantidad, sin encabezado, desde la fila 7

Se carga TODO lo pedido aunque no haya stock: el sistema avisa las lineas no
satisfechas y eso lo resuelve el deposito. El disponible se informa igual.

Controles: cantidades numericas y positivas, digito verificador del ISBN,
ISBN existente en SII/LP, CODIGO (ID) del archivo contra el del sistema, ISBN
repetido en un archivo, cuenta del nombre contra la de adentro, cliente en el
maestro (activo, del vendedor, con direccion de entrega), y cuadre de unidades
entre los archivos de entrada y los de salida.

Uso:
    python ampliaciones.py CARPETA --vendedor MARTIRENA --salida ./out \
        --sii SII.xls --lp LP.xlsx --maestro Maestro_clientes.xls \
        [--reemplazar 9789500206174=9789500209052] [--excluir 21896] [--extra archivo.xls]
"""
import argparse
import glob
import os
import re
import sys
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).parent))
from isbn import resolver  # noqa: E402

PATRON = re.compile(r'AMPLIACION\s*-\s*(.+?)\s+(\d{4,6})\s+\d{2}-\d{2}-\d{4}', re.I)
_NS = {'table': 'urn:oasis:names:tc:opendocument:xmlns:table:1.0',
       'office': 'urn:oasis:names:tc:opendocument:xmlns:office:1.0',
       'text': 'urn:oasis:names:tc:opendocument:xmlns:text:1.0'}


def _q(p, t):
    return '{%s}%s' % (_NS[p], t)


def leer_ods(ruta):
    """Primera hoja de un .ods como lista de filas (odfpy no siempre instala)."""
    root = ET.fromstring(zipfile.ZipFile(ruta).read('content.xml'))
    tabla = next(root.iter(_q('table', 'table')))
    filas = []
    for row in tabla.iter(_q('table', 'table-row')):
        fila = []
        for cell in row:
            if cell.tag not in (_q('table', 'table-cell'), _q('table', 'covered-table-cell')):
                continue
            n = min(int(cell.get(_q('table', 'number-columns-repeated'), '1')), 50)
            vt = cell.get(_q('office', 'value-type'))
            if vt in ('float', 'currency', 'percentage'):
                v = float(cell.get(_q('office', 'value')))
            elif vt is None:
                v = None
            else:
                v = ' '.join(''.join(p.itertext()) for p in cell.findall(_q('text', 'p')))
            fila += [v] * n
        while fila and fila[-1] is None:
            fila.pop()
        filas.append(fila)
    while filas and not filas[-1]:
        filas.pop()
    return filas


def vacio(v):
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == ''


def nombre_archivo(ruta):
    base = re.sub(r'^[0-9a-f]{8}-', '', Path(ruta).name).replace('_', ' ')
    base = unicodedata.normalize('NFC', base.replace('#U0303', '̃'))
    m = PATRON.search(base)
    if not m:
        raise SystemExit(f'{base}: el nombre no sigue "AMPLIACION - CLIENTE CUENTA dd-mm-aaaa"')
    return base, m.group(1).strip(), int(m.group(2))


def leer(ruta):
    base, cliente, cuenta = nombre_archivo(ruta)
    if ruta.lower().endswith('.ods'):
        filas = leer_ods(ruta)
    else:
        filas = pd.read_excel(ruta, sheet_name=0, header=None, dtype=object).values.tolist()
    cab = [str(c).strip().upper() for c in filas[0]]
    lineas = []
    for k, r in enumerate(filas[1:], start=2):
        d = dict(zip(cab, r))
        # CODIGO es el ID cuando hay columna ISBN; si no, CODIGO trae el ISBN
        isbn_raw = d.get('ISBN') if 'ISBN' in d else d.get('CODIGO')
        cant = d.get('AMPLIACION')
        if vacio(isbn_raw) and vacio(cant):
            continue  # filas de relleno (cuenta y cliente sin titulo)
        lineas.append({'archivo': base, 'cliente': cliente, 'cuenta': cuenta, 'fila': k,
                       'cuenta_dentro': d.get('CUENTA'),
                       'id': d.get('CODIGO') if 'ISBN' in d else None,
                       'isbn_raw': isbn_raw, 'titulo': d.get('TITULO'), 'cant': cant})
    return lineas


def a_int(v):
    try:
        return int(float(str(v).strip()))
    except (TypeError, ValueError):
        return None


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('carpeta')
    ap.add_argument('--vendedor', required=True)
    ap.add_argument('--salida', default='./ampliaciones')
    ap.add_argument('--sii')
    ap.add_argument('--lp')
    ap.add_argument('--maestro')
    ap.add_argument('--reemplazar', action='append', default=[], help='ISBN_VIEJO=ISBN_NUEVO')
    ap.add_argument('--excluir', action='append', default=[], type=int, help='cuenta a dejar afuera')
    ap.add_argument('--extra', action='append', default=[],
                    help='archivo que reemplaza al de la misma cuenta en la carpeta')
    args = ap.parse_args()

    rutas = sorted(glob.glob(os.path.join(args.carpeta, 'AMPLIACION*')))
    extra = {nombre_archivo(r)[2]: r for r in args.extra}
    rutas = [r for r in rutas if nombre_archivo(r)[2] not in extra] + list(extra.values())
    rutas = [r for r in rutas if nombre_archivo(r)[2] not in args.excluir]
    problemas, avisos = [], []
    for c, r in extra.items():
        avisos.append(f'cuenta {c}: se usa {Path(r).name} en lugar del de la carpeta')

    df = pd.DataFrame([l for r in rutas for l in leer(r)])
    df['q'] = df.cant.map(lambda v: pd.to_numeric(v, errors='coerce'))
    for _, l in df[df.q.isna() | (df.q <= 0) | (df.q != df.q.round())].iterrows():
        problemas.append(f"{l.archivo} fila {l.fila}: cantidad {l.cant!r}")

    def norm(v):
        if vacio(v):
            return None, 'vacio'
        return resolver(str(int(v)) if isinstance(v, float) else str(v).strip())
    df[['isbn', 'metodo']] = df.isbn_raw.apply(lambda v: pd.Series(norm(v)))
    for _, l in df[df.metodo != 'ya era ISBN-13'].iterrows():
        problemas.append(f"{l.archivo} fila {l.fila}: ISBN {l.isbn_raw!r} -> {l.isbn} ({l.metodo})")
    reemplazos = dict(r.split('=') for r in args.reemplazar)
    for viejo, nuevo in reemplazos.items():
        n = (df.isbn == viejo).sum()
        avisos.append(f'reemplazo {viejo} -> {nuevo} en {n} linea(s)')
        df.loc[df.isbn == viejo, 'isbn'] = nuevo

    sistema = {}
    disp = {}
    if args.sii:
        s = pd.read_excel(args.sii)
        s = s[s['ID ARTICULO'].notna()]
        for _, r in s.iterrows():
            k = a_int(r['ISBN'])
            if k:
                sistema[str(k)] = (int(r['ID ARTICULO']), str(r['TITULO']).strip())
                disp[str(k)] = int(r['DISPONIBLE'] or 0)
    if args.lp:
        lp = pd.read_excel(args.lp, dtype={'Código / ISBN': str})
        for _, r in lp.iterrows():
            k = str(r['Código / ISBN']).replace('-', '')
            sistema.setdefault(k, (int(r['Nro. Art.']), str(r['Título']).strip()))
    if sistema:
        for _, l in df[~df.isbn.isin(sistema.keys())].iterrows():
            problemas.append(f"{l.archivo} fila {l.fila}: ISBN {l.isbn} ({l.titulo}) no esta en SII ni LP")
        for _, l in df[df.id.map(a_int).notna()].iterrows():
            if l.isbn in sistema and a_int(l.id) != sistema[l.isbn][0]:
                problemas.append(f"{l.archivo} fila {l.fila}: CODIGO {a_int(l.id)} pero el ISBN "
                                 f"{l.isbn} es el ID {sistema[l.isbn][0]} ({sistema[l.isbn][1]})")
    for _, l in df[df.duplicated(['archivo', 'isbn'], keep=False)].iterrows():
        problemas.append(f"{l.archivo} fila {l.fila}: ISBN {l.isbn} repetido en el archivo")
    for _, l in df[df.cuenta_dentro.map(a_int).notna()
                   & (df.cuenta_dentro.map(a_int) != df.cuenta)].iterrows():
        avisos.append(f"{l.archivo} fila {l.fila}: la columna CUENTA dice {l.cuenta_dentro!r} "
                      f"(el nombre dice {l.cuenta}; no afecta la carga)")
    for _, l in df[df.cuenta_dentro.notna() & df.cuenta_dentro.map(a_int).isna()
                   & ~df.cuenta_dentro.map(vacio)].iterrows():
        avisos.append(f"{l.archivo} fila {l.fila}: la columna CUENTA dice {l.cuenta_dentro!r} "
                      f"(no afecta la carga)")

    if args.maestro:
        m = pd.read_excel(args.maestro)
        for cuenta, cliente in df[['cuenta', 'cliente']].drop_duplicates().values:
            mm = m[m.IDCLIENTE == cuenta]
            if mm.empty:
                problemas.append(f'cuenta {cuenta} ({cliente}) no esta en el maestro')
                continue
            if mm.FECHABAJA.notna().all():
                problemas.append(f'cuenta {cuenta} ({mm.iloc[0].NOMBRE}) dada de baja')
            if (mm['TIPO DIRECCION'] == 'ENTR').sum() == 0:
                problemas.append(f'cuenta {cuenta} ({mm.iloc[0].NOMBRE}) sin direccion de entrega')
            if args.vendedor.upper() not in unicodedata.normalize('NFKD', str(mm.iloc[0].VENDEDOR)) \
                    .encode('ascii', 'ignore').decode().upper():
                avisos.append(f'cuenta {cuenta} ({mm.iloc[0].NOMBRE}) es de "{mm.iloc[0].VENDEDOR}"')

    # --- archivos de carga ---
    out = Path(args.salida)
    out.mkdir(parents=True, exist_ok=True)
    escritas = 0
    for (cuenta, cliente), g in df.groupby(['cuenta', 'cliente'], sort=False):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Hoja1'
        for i, l in enumerate(g.itertuples(), start=7):
            ws.cell(i, 1, int(l.isbn)).number_format = '0'
            ws.cell(i, 2, int(l.q))
        ws.column_dimensions['A'].width = 16
        slug = re.sub(r'\s+', '_', cliente.upper())
        wb.save(out / f'{args.vendedor.upper()}_-_{slug}_-_{cuenta}.xlsx')
        escritas += 1

    # cuadre por otro camino: releer lo escrito
    releido = 0
    for f in out.glob(f'{args.vendedor.upper()}_-_*.xlsx'):
        ws = load_workbook(f).active
        releido += sum(r[1] for r in ws.iter_rows(min_row=7, values_only=True) if r[1])

    print(f'{len(rutas)} archivos | {df.cuenta.nunique()} clientes | {len(df)} lineas | '
          f'{int(df.q.sum())} u | {df.isbn.nunique()} ISBN')
    print(f'escritos {escritas} archivos en {out}/ | cuadre entrada {int(df.q.sum())} = '
          f'releido {releido} {"OK" if releido == int(df.q.sum()) else "*** NO CUADRA ***"}')
    if disp:
        g = df.groupby('isbn').q.sum()
        cortos = [(i, int(q), disp.get(i)) for i, q in g.items() if disp.get(i) is None or q > disp[i]]
        print(f'\ntitulos donde lo pedido supera el disponible del SII ({len(cortos)}; se cargan igual):')
        for i, q, d in sorted(cortos, key=lambda x: -x[1]):
            print(f'  {i}  pide {q:>4}  disp {d if d is not None else "-":>5}  '
                  f'{sistema.get(i, ("", ""))[1][:40]}')
    print(f'\n{len(problemas)} PROBLEMAS')
    for p in problemas:
        print('  *** ' + p)
    print(f'\n{len(avisos)} avisos')
    for a in avisos:
        print('  -   ' + a)
    df.drop(columns=['q']).to_csv(out / 'detalle_lineas.csv', sep=';', index=False, encoding='utf-8-sig')


if __name__ == '__main__':
    main()
