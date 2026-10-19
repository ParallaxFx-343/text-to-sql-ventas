"""Pasa a Excel una factura de GRUPO ILHSA en PDF (texto extraido con pdfplumber).
Uso: python -P factura_xlsx.py FACTURA.pdf SALIDA.xlsx"""
import sys, re, pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
fin, fout = sys.argv[1:3]
num = lambda s: float(s.replace('.', '').replace(',', '.'))
pat = re.compile(r'^(\d{13})(.+?) (\d+) ([\d.]+,\d{2}) (\d+) (\d+) ([\d.]+,\d{2}) ([\d.]+,\d{2})$')
items, subt, txt = [], [], []
with pdfplumber.open(fin) as pdf:
    for pg in pdf.pages:
        t = pg.extract_text(); txt.append(t); n_pag = 0
        for l in t.split('\n'):
            m = pat.match(l.strip())
            if m:
                i, tit, q, pu, iva, dto, desc, imp = m.groups()
                items.append([int(i), tit.strip(), int(q), num(pu), int(iva), int(dto), num(desc), num(imp)])
            ms = re.search(r'SUBTOTAL \$: ([\d.]+,\d{2})', l)
            if ms: subt.append((len(items), num(ms.group(1))))
full = '\n'.join(txt)
g = lambda p: (re.search(p, full).group(1).strip() if re.search(p, full) else '')
cab = [('Factura', 'B ' + g(r'Nº (\d{4} - \d{8})')), ('Fecha', g(r'Fecha: (\d{2}/\d{2}/\d{4})')),
       ('Cliente', g(r'Cliente:(.+?) C\.U\.I\.T\.')), ('CUIT cliente', g(r'C\.U\.I\.T\.: (\d{11})')),
       ('Vendedor', g(r'VENDEDOR: (.+?) NOTA')), ('Nota', g(r'NOTA: (\S+)')),
       ('Lista de precios', g(r'LISTA DE PRECIOS: (.+?) Tipo de Compra')), ('Vencimiento', g(r'Vencimiento: (\d{2}/\d{2}/\d{4})')),
       ('C.A.E.', g(r'C\.A\.E\. Nº:(\d+)')), ('Vto. C.A.E.', g(r'FECHA DE VTO\.: (\d{2}/\d{2}/\d{4})'))]
total = num(g(r'(?<!SUB)TOTAL \$: ([\d.]+,\d{2})'))
# controles
for it in items:
    assert abs(it[2] * it[3] * it[5] / 100 - it[6]) < 0.01 and abs(it[2] * it[3] - it[6] - it[7]) < 0.01, it
acum = 0; k0 = 0
for k, s in subt:
    acum = sum(x[7] for x in items[:k]); assert abs(acum - s) < 0.01, (k, acum, s)
assert abs(sum(x[7] for x in items) - total) < 0.01
wb = Workbook(); ws = wb.active; ws.title = 'Factura'
for r, (a, b) in enumerate(cab, 1):
    ws.cell(r, 1, a).font = Font(bold=True); ws.cell(r, 2, b)
h0 = len(cab) + 2
cols = ['ISBN', 'Título', 'Cantidad', 'Precio unitario', 'Alícuota IVA', '% Dto', 'Descuento', 'Importe neto']
fino = Side(style='thin', color='D9DCE1'); caja = Border(left=fino, right=fino, top=fino, bottom=fino)
for j, c in enumerate(cols, 1):
    x = ws.cell(h0, j, c); x.font = Font(bold=True, color='FFFFFF'); x.fill = PatternFill('solid', fgColor='1F4E78')
    x.alignment = Alignment(horizontal='center'); x.border = caja
for i, it in enumerate(items, h0 + 1):
    for j, v in enumerate(it, 1):
        x = ws.cell(i, j, v); x.border = caja
    ws.cell(i, 1).number_format = '0'
    for j in (4, 7, 8): ws.cell(i, j).number_format = '#,##0.00'
u = h0 + len(items); t = u + 1
ws.cell(t, 2, 'TOTAL').font = Font(bold=True)
for j, L in ((3, 'C'), (7, 'G'), (8, 'H')):
    x = ws.cell(t, j, f'=SUM({L}{h0+1}:{L}{u})'); x.font = Font(bold=True); x.border = caja
    if j > 3: x.number_format = '#,##0.00'
for L, w in zip('ABCDEFGH', (16, 52, 10, 15, 12, 8, 14, 15)): ws.column_dimensions[L].width = w
ws.freeze_panes = ws.cell(h0 + 1, 1); ws.auto_filter.ref = f'A{h0}:H{u}'
wb.save(fout)
print(f'{len(items)} lineas | {sum(x[2] for x in items)} ejemplares | neto {sum(x[7] for x in items):,.2f} = TOTAL factura {total:,.2f} | subtotales por pagina OK ({len(subt)})')
print(cab)
